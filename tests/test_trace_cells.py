# trace_cells 의 셀 분류와 참조 추적 시험.
#   uv run --no-project --with openpyxl --with pytest pytest tests/test_trace_cells.py
import sys
from pathlib import Path

import openpyxl
from openpyxl.workbook.defined_name import DefinedName

sys.path.insert(0, str(Path(__file__).resolve().parents[1]
                       / "plugins/kw-dashboard/skills/explaining-excel-cell/scripts"))
import trace_cells  # noqa: E402


def make_book(tmp_path, sheets, names=None):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for title, cells in sheets.items():
        ws = wb.create_sheet(title)
        for coord, value in cells.items():
            ws[coord] = value
    for name, target in (names or {}).items():
        wb.defined_names[name] = DefinedName(name, attr_text=target)
    path = tmp_path / "book.xlsx"
    wb.save(path)
    return path


def run(tmp_path, sheets, results, names=None):
    return trace_cells.trace(make_book(tmp_path, sheets, names), results)


def test_result_cells_are_classified_by_kind(tmp_path):
    t = run(tmp_path, {"R": {"A1": "제목", "B1": 3, "C1": "=B1*2"}}, ["R"])
    assert t.cells["R"][(1, 1)].kind == "value"
    assert t.cells["R"][(1, 2)].kind == "value"
    assert t.cells["R"][(1, 3)].kind == "formula"
    assert all(c.reached for c in t.cells["R"].values())


def test_only_referenced_cells_of_other_sheets_are_reached(tmp_path):
    t = run(tmp_path, {"R": {"A1": "=S!A1+S!B1"}, "S": {"A1": 1, "B1": 2, "C1": 3}}, ["R"])
    s = t.cells["S"]
    assert s[(1, 1)].reached and s[(1, 2)].reached
    assert not s[(1, 3)].reached


def test_tracing_follows_formula_chains(tmp_path):
    t = run(tmp_path, {"R": {"A1": "=S!A1"}, "S": {"A1": "=SUM(T!A1:A2)"},
                       "T": {"A1": 1, "A2": 2, "A3": 9}}, ["R"])
    assert t.cells["S"][(1, 1)].reached
    assert t.cells["T"][(1, 1)].reached and t.cells["T"][(2, 1)].reached
    assert not t.cells["T"][(3, 1)].reached


def test_whole_column_range_reaches_existing_cells_only(tmp_path):
    t = run(tmp_path, {"R": {"A1": "=VLOOKUP(1,S!$A:$B,2,FALSE)"},
                       "S": {"A1": 1, "B1": "x", "A5": 2, "C1": "no"}}, ["R"])
    s = t.cells["S"]
    assert s[(1, 1)].reached and s[(1, 2)].reached and s[(5, 1)].reached
    assert not s[(1, 3)].reached


def test_quoted_sheet_name_with_space(tmp_path):
    t = run(tmp_path, {"R": {"A1": "='2.설정액 변환용'!B2"}, "2.설정액 변환용": {"B2": 5}}, ["R"])
    assert t.cells["2.설정액 변환용"][(2, 2)].reached


def test_unquoted_korean_sheet_name(tmp_path):
    t = run(tmp_path, {"R": {"A1": "=VLOOKUP(A2,운용펀드_리스트!$C:$D,2,0)"},
                       "운용펀드_리스트": {"C1": "k", "D1": "v"}}, ["R"])
    assert t.cells["운용펀드_리스트"][(1, 4)].reached


def test_external_reference_is_recorded_not_traced(tmp_path):
    t = run(tmp_path, {"R": {"A1": "=[1]Sheet1!A1"}}, ["R"])
    assert t.external_refs == {("R", (1, 1)): ["[1]Sheet1!A1"]}


def test_defined_name_is_resolved(tmp_path):
    t = run(tmp_path, {"R": {"A1": "=SUM(Rates)"}, "S": {"A1": 1, "A2": 2, "A3": 3}}, ["R"],
            names={"Rates": "S!$A$1:$A$2"})
    assert t.cells["S"][(1, 1)].reached and t.cells["S"][(2, 1)].reached
    assert not t.cells["S"][(3, 1)].reached


def test_dynamic_reference_is_flagged(tmp_path):
    t = run(tmp_path, {"R": {"A1": '=INDIRECT("S!A1")'}, "S": {"A1": 1}}, ["R"])
    assert ("R", (1, 1)) in t.dynamic_refs


def test_missing_sheet_and_ref_error_are_flagged(tmp_path):
    t = run(tmp_path, {"R": {"A1": "=Gone!A1", "A2": "=#REF!A1+1"}}, ["R"])
    assert ("R", (1, 1)) in t.broken_refs
    assert ("R", (2, 1)) in t.broken_refs


def make_table_book(tmp_path, formula_cells):
    from openpyxl.worksheet.table import Table
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "D"
    for coord, v in {"A1": "일자", "B1": "값", "C1": "계산",
                     "A2": 1, "B2": 10, "A3": 2, "B3": 20}.items():
        ws[coord] = v
    for coord, f in formula_cells.items():
        ws[coord] = f
    ws.add_table(Table(displayName="T1", ref="A1:C3"))
    r = wb.create_sheet("R")
    r["A1"] = "=SUM(T1[값])"
    path = tmp_path / "table.xlsx"
    wb.save(path)
    return path


def test_structured_reference_to_column_reaches_data_rows(tmp_path):
    t = trace_cells.trace(make_table_book(tmp_path, {}), ["R"])
    d = t.cells["D"]
    assert d[(2, 2)].reached and d[(3, 2)].reached
    assert not d[(1, 2)].reached and not d[(2, 1)].reached


def test_structured_this_row_reference(tmp_path):
    path = make_table_book(tmp_path, {"C2": "=T1[[#This Row],[일자]]*2", "C3": "=[@일자]*2"})
    t = trace_cells.trace(path, ["D"])
    assert not t.broken_refs


def test_pivot_output_depends_on_its_source(tmp_path):
    import zipfile
    path = make_book(tmp_path, {"Src": {"A1": "k", "B1": "v", "A2": "x", "B2": 1, "C2": "no"},
                                "P": {"A1": "x", "B1": 1}})
    pivot = ('<pivotTableDefinition xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" '
             'name="p" cacheId="1"><location ref="A1:B1" firstHeaderRow="0" firstDataRow="0" '
             'firstDataCol="0"/></pivotTableDefinition>')
    cache = ('<pivotCacheDefinition xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
             '<cacheSource type="worksheet"><worksheetSource ref="A1:B2" sheet="Src"/></cacheSource>'
             '</pivotCacheDefinition>')
    rel = ('<?xml version="1.0"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/'
           'relationships"><Relationship Id="rId9" Type="{t}" Target="{target}"/></Relationships>')
    with zipfile.ZipFile(path, "a") as z:
        z.writestr("xl/pivotTables/pivotTable1.xml", pivot)
        z.writestr("xl/pivotTables/_rels/pivotTable1.xml.rels",
                   rel.format(t="pivotCacheDefinition", target="../pivotCache/pivotCacheDefinition1.xml"))
        z.writestr("xl/pivotCache/pivotCacheDefinition1.xml", cache)
        z.writestr("xl/worksheets/_rels/sheet2.xml.rels",
                   rel.format(t="pivotTable", target="../pivotTables/pivotTable1.xml"))
    t = trace_cells.trace(path, ["P"])
    assert t.cells["P"][(1, 2)].kind == "pivot"
    assert t.cells["Src"][(2, 2)].reached
    assert not t.cells["Src"][(2, 3)].reached


def test_precedents_are_recorded_per_formula(tmp_path):
    t = run(tmp_path, {"R": {"A1": "=S!A1+SUM(S!B1:B3)"}, "S": {"A1": 1, "B1": 2}}, ["R"])
    assert t.precedents[("R", (1, 1))] == [("S", (1, 1, 1, 1)), ("S", (1, 2, 3, 2))]


def test_reference_tokens_keep_formula_text_and_areas(tmp_path):
    t = run(tmp_path, {"R": {"A1": "=S!A1+SUM(S!B1:B3)+[1]X!A1"}, "S": {"A1": 1, "B1": 2}}, ["R"])
    assert t.ref_tokens[("R", (1, 1))] == [
        ("S!A1", [("S", (1, 1, 1, 1))]), ("S!B1:B3", [("S", (1, 2, 3, 2))]), ("[1]X!A1", None)]


def test_page_has_a_tab_per_used_sheet_and_cell_values(tmp_path):
    import json
    import re
    path = make_book(tmp_path, {"R": {"A1": "=S!A1*2"}, "S": {"A1": 21}, "X": {"A1": 0}})
    t = trace_cells.trace(path, ["R"])
    page = trace_cells.render_html(t, trace_cells.load_values(path, ["R", "S"]), "book")
    data = json.loads(re.search(r'<script id="data" type="application/json">(.*?)</script>',
                                page, re.S).group(1))
    assert [s["name"] for s in data["sheets"]] == ["R", "S"]
    assert data["unused_sheets"] == ["X"]
    s_cells = data["sheets"][1]["cells"]
    assert s_cells[0][:4] == [1, 1, "r", "21"]
    r_cell = data["sheets"][0]["cells"][0]
    assert r_cell[4] == "=S!A1*2" and r_cell[5] == [[1, 1, 1, 1, 1]]


def test_vlookup_reaches_only_key_and_return_columns(tmp_path):
    t = run(tmp_path, {"R": {"A1": "=VLOOKUP(1,S!$A:$C,3,FALSE)"},
                       "S": {"A1": 1, "B1": "unused", "C1": "v"}}, ["R"])
    s = t.cells["S"]
    assert s[(1, 1)].reached and s[(1, 3)].reached
    assert not s[(1, 2)].reached


def test_hlookup_reaches_only_key_and_return_rows(tmp_path):
    t = run(tmp_path, {"R": {"A1": "=HLOOKUP(1,S!A1:C3,3,FALSE)"},
                       "S": {"A1": 1, "A2": "unused", "A3": "v"}}, ["R"])
    s = t.cells["S"]
    assert s[(1, 1)].reached and s[(3, 1)].reached
    assert not s[(2, 1)].reached


def test_vlookup_with_computed_column_keeps_whole_range(tmp_path):
    t = run(tmp_path, {"R": {"A1": "=VLOOKUP(1,S!$A:$C,R!B1,FALSE)", "B1": 2},
                       "S": {"A1": 1, "B1": "x", "C1": "v"}}, ["R"])
    assert all(c.reached for c in t.cells["S"].values())


def test_unknown_result_sheet_raises(tmp_path):
    path = make_book(tmp_path, {"R": {"A1": 1}})
    try:
        trace_cells.trace(path, ["없는시트"])
    except ValueError as e:
        assert "없는시트" in str(e)
    else:
        raise AssertionError("ValueError expected")


def test_display_handles_date_datetime_and_time():
    import datetime as dt
    assert trace_cells._display(dt.date(2026, 9, 30)) == "2026-09-30"
    assert trace_cells._display(dt.datetime(2026, 9, 30)) == "2026-09-30"
    assert trace_cells._display(dt.datetime(2026, 9, 30, 9, 5)) == "2026-09-30 09:05:00"
    assert trace_cells._display(dt.time(9, 5)) == "09:05:00"


def test_make_resolver_resolves_structured_external_and_plain_refs():
    table = trace_cells.Table("S", (5, 1, 7, 2), ["키", "값"], 1, 0)
    resolve = trace_cells.make_resolver({"S": {}, "T": {}}, {"T1": table}, {})
    assert resolve("T1[[#This Row],[값]]", "S", (6, 3)) == [("S", (6, 2, 6, 2))]
    assert resolve("T!$A$1", "S", (1, 1)) == [("T", (1, 1, 1, 1))]
    assert resolve("'[1]X'!A1", "S", (1, 1)) == "external"
    assert resolve("없는시트!A1", "S", (1, 1)) == "broken"

# capture_views 시험. openpyxl 로 만든 작은 엑셀로 덩어리 나누기를 본다.
#   uv run --no-project --with openpyxl --with pytest pytest tests/test_capture_views.py
import os
import sys
from pathlib import Path

import openpyxl
import pytest
from openpyxl.worksheet.table import Table

ROOT = Path(__file__).resolve().parents[1] / "plugins/kw-dashboard/skills"
sys.path.insert(0, str(ROOT / "explaining-excel-cell/scripts"))
import capture_views  # noqa: E402
import explain_cell  # noqa: E402


def make_book(tmp_path, cells, table=None):
    """cells 는 {"시트": {"A1": 값}}, table 은 ("시트", "표이름", "A5:C8") 이다."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for sheet, values in cells.items():
        ws = wb.create_sheet(sheet)
        for coord, v in values.items():
            ws[coord] = v
    if table:
        sheet, name, ref = table
        wb[sheet].add_table(Table(displayName=name, ref=ref))
    path = tmp_path / "b.xlsx"
    wb.save(path)
    return explain_cell.Book(str(path))


TABLE = {"A5": "일자", "B5": "금액", "C5": "구분",
         "A6": 1, "B6": 10, "C6": "x",
         "A7": 2, "B7": 20, "C7": "y",
         "A8": 3, "B8": 30, "C8": "z"}


def test_table_block_starts_at_title_band_row(tmp_path):
    book = make_book(tmp_path, {"S": {**TABLE, "A3": "펀드마스타"}}, ("S", "T1", "A5:C8"))
    for r, c in [(5, 1), (7, 3), (8, 2)]:  # 표의 어느 칸을 줘도 같은 덩어리다
        blk = capture_views.block_of(book, "S", r, c)
        assert blk["rect"] == (3, 1, 8, 3)
        assert blk["table"] == "T1"


def test_table_block_starts_at_header_row_without_title_band(tmp_path):
    book = make_book(tmp_path, {"S": TABLE}, ("S", "T1", "A5:C8"))
    assert capture_views.block_of(book, "S", 6, 2)["rect"] == (5, 1, 8, 3)


def test_blocks_split_by_one_empty_column_and_one_empty_row(tmp_path):
    cells = {"S": {"A1": 1, "B1": 2, "D1": 3, "E1": 4,      # 빈 열 C 가 가른다
                   "A4": 5, "B4": 6}}                       # 빈 행 2 · 3 이 가른다
    book = make_book(tmp_path, cells)
    assert sorted(capture_views.detect_blocks(book, "S")) == [(1, 1, 1, 2), (1, 4, 1, 5), (4, 1, 4, 2)]


def test_cell_outside_any_block_is_its_own_block(tmp_path):
    book = make_book(tmp_path, {"S": {"A1": 1}})
    assert capture_views.block_of(book, "S", 9, 9)["rect"] == (9, 9, 9, 9)


def test_label_skips_error_text_and_names_range(tmp_path):
    book = make_book(tmp_path, {"S": {"A1": "#REF!", "A2": "일자", "B2": 1, "A3": 5}})
    blk = capture_views.block_of(book, "S", 2, 1)
    assert capture_views.label_of(book, blk) == "S · 일자 덩어리 A1:B3"


def test_label_of_table_says_table(tmp_path):
    book = make_book(tmp_path, {"S": {**TABLE, "A3": "펀드마스타"}}, ("S", "T1", "A5:C8"))
    blk = capture_views.block_of(book, "S", 6, 1)
    assert capture_views.label_of(book, blk) == "S · 펀드마스타 표 A3:C8"


def test_detect_blocks_runs_once_per_sheet(tmp_path, monkeypatch):
    book = make_book(tmp_path, {"S": {"A1": 1, "B2": 2}})
    calls = []
    real = capture_views.detect_blocks
    monkeypatch.setattr(capture_views, "detect_blocks", lambda b, s: calls.append(s) or real(b, s))
    for _ in range(100):
        capture_views.block_of(book, "S", 1, 1)
    assert calls == ["S"]


def views_of(tmp_path, cells, target, sheet="S"):
    book = make_book(tmp_path, cells)
    data = explain_cell.build(book, sheet, target)
    return data, capture_views.build_views(book, data["steps"], data["repeats"])


def test_candidates_hold_head_rows_and_rows_around_a_far_cell(tmp_path):
    cells = {"S": {**{f"A{r}": r for r in range(1, 121)}, "B100": "=A100*2"}}
    data, views = views_of(tmp_path, cells, "B100")
    v = views[-1]
    assert v["cur"] == [100, 2] and v["label"] == "S · 덩어리 A1:B120"
    assert set(range(1, 25)) <= set(v["rows"])          # 덩어리 머리 행
    assert set(range(94, 107)) <= set(v["rows"])        # 가리키는 칸 둘레
    assert 50 not in v["rows"]                          # 가운데는 후보가 아니다
    assert v["cells"]["100,1"] == "100"


def test_input_in_same_block_goes_to_inputs_and_other_block_to_far(tmp_path):
    cells = {"S": {"A1": 5, "B1": "=A1*2", "E1": "=A1+1"}}  # 빈 열 C · D 가 E1 을 다른 덩어리로 가른다
    data, views = views_of(tmp_path, cells, "E1")
    v = views[-1]
    assert v["inputs"] == []
    assert v["far"] == [{"j": 0, "addr": "S!A1", "value": "5", "label": "S · 덩어리 A1:B1"}]
    data, views = views_of(tmp_path, cells, "B1")
    assert views[-1]["inputs"] == [[1, 1]] and views[-1]["far"] == []


def test_cell_text_stays_raw(tmp_path):
    cells = {"S": {"A1": "<img src=x onerror=alert(1)>", "B1": '=A1&"x"'}}
    data, views = views_of(tmp_path, cells, "B1")
    assert views[-1]["cells"]["1,1"] == "<img src=x onerror=alert(1)>"


CHAIN = {"S": {"A1": 1.0, **{f"A{r}": f"=A{r - 1}+1" for r in range(2, 21)}}}


def test_folded_range_has_one_view_with_first_cell_as_input_and_last_cell_as_current(tmp_path):
    data, views = views_of(tmp_path, CHAIN, "A20")
    rep = data["repeats"][0]
    assert len(views) == len(data["steps"])
    assert views[rep["first"]]["inputs"] == [[5, 1]] and views[rep["first"]]["cur"] == [19, 1]
    assert all(views[i] is None for i in range(rep["first"] + 1, rep["last"] + 1))
    assert views[rep["last"] + 1]["cur"] == [20, 1]     # 접힌 구간 뒤 단계는 보통 항목이다


def test_earlier_keeps_only_the_latest_20_and_skips_folded_steps(tmp_path):
    cells = {"S": {f"A{r}": r for r in range(1, 41)}}
    cells["S"]["B1"] = "=SUM(A1:A40)"
    data, views = views_of(tmp_path, cells, "B1")
    assert all(len(v["earlier"]) <= 20 for v in views if v)


def test_long_chain_views_stay_small_and_fast(tmp_path):
    import json
    import time
    cells = {"S": {"A1": 1.0, **{f"A{r}": f"=A{r - 1}+1" for r in range(2, 1723)}}}
    book = make_book(tmp_path, cells)
    data = explain_cell.build(book, "S", "A1722")
    t = time.time()
    views = capture_views.build_views(book, data["steps"], data["repeats"])
    assert time.time() - t <= 10
    assert len(json.dumps(views, ensure_ascii=False)) <= 5_000_000


def test_build_sources_groups_value_steps_by_kind_in_fixed_order():
    nodes = [{"id": "n1", "title": "통화"}, {"id": "n2", "title": "금액"}, {"id": "n3", "title": "계산"}]
    steps = [
        {"node": "n1", "kind": "value", "source": "수기 입력", "addr": "S!A1", "value": "KRW"},
        {"node": "n2", "kind": "import", "source": "WFS", "addr": "S!B1", "value": ""},
        {"node": "n2", "kind": "value", "source": "WFS", "addr": "S!B2", "value": "10"},
        {"node": "n2", "kind": "value", "source": "WINUS", "addr": "S!B3", "value": "20"},
        {"node": "n3", "kind": "calc", "source": "계산", "addr": "S!C1", "value": "30"},
        {"node": "n1", "kind": "value", "source": "피벗", "addr": "S!A2", "value": "1"},
    ]
    assert capture_views.build_sources(nodes, steps) == [
        {"kind": "WINUS", "name": "WINUS", "items": [{"j": 3, "name": "금액", "value": "20"}]},
        {"kind": "WFS", "name": "WFS(우리펀드 서비스)", "items": [{"j": 2, "name": "금액", "value": "10"}]},
        {"kind": "수기 입력", "name": "수기 입력", "items": [{"j": 0, "name": "통화", "value": "KRW"}]},
    ]


def test_build_returns_views_and_sources_with_manual_and_external_file_items(tmp_path):
    cells = {"S": {"A1": 3.0, "B1": "=[1]Liquidity!C5*2", "D1": 4.0, "E1": "=A1+B1+D1"}}
    data = explain_cell.build(make_book(tmp_path, cells), "S", "E1")
    assert len(data["views"]) == len(data["steps"])
    kinds = {s["kind"]: s for s in data["sources"]}
    assert [len(kinds[k]["items"]) for k in ("외부 파일", "수기 입력")] == [1, 2]
    assert all(data["steps"][it["j"]]["kind"] == "value" for s in data["sources"] for it in s["items"])


def test_wfs_connection_appears_as_wfs_source(tmp_path):
    book = make_book(tmp_path, {"S": {"A3": "펀드코드", "A4": "2043", "B3": "설정", "B4": 5.0, "D1": "=A4&B4"}})
    book.tables["표_Q"] = explain_cell.trace_cells.Table("S", (3, 1, 4, 2), ["펀드코드", "설정"], 1, 0, query=True)
    book.queries = {"표_Q": {"name": "q", "source": "WFS", "tables": ["T1@TESTLINK3"]}}
    book.loaded_ranges = [{"sheet": "S", "area": (3, 1, 4, 2), "source": "WFS", "kind": "ODBC 쿼리 테이블", "name": "q",
                           "tables": ["T1@TESTLINK3"]}]
    sources = explain_cell.build(book, "S", "D1")["sources"]
    assert [s["name"] for s in sources if s["kind"] == "WFS"] == ["WFS(우리펀드 서비스)"]
    assert {it["name"] for s in sources if s["kind"] == "WFS" for it in s["items"]} == {"펀드코드", "설정"}


FUNDINFO = Path(os.environ.get("EXPLAIN_FUNDINFO_XLSX", "/없음"))


@pytest.mark.skipif(not FUNDINFO.exists(), reason="EXPLAIN_FUNDINFO_XLSX 가 가리키는 표본 엑셀이 없다")
def test_fundinfo_g2_lists_manual_input_value_name():
    sources = explain_cell.build(FUNDINFO, "4.수치 수합", "G2")["sources"]
    manual = [it["name"] for s in sources if s["kind"] == "수기 입력" for it in s["items"]]
    assert "실현수익률-1" in manual


def test_fold_input_outside_the_block_is_not_drawn_inside_it(tmp_path):
    values = {f"B{r}": float(r) for r in (2, 3, 4, 5, 6, 7, 9, 10)}  # 빈 행 8 이 두 덩어리로 가른다
    cells = {"S": {**values, "D1": "=" + "+".join(f"B{r}" for r in (2, 3, 4, 5, 6, 7, 9, 10))}}
    data, views = views_of(tmp_path, cells, "D1")
    rep = data["repeats"][0]
    v = views[rep["first"]]
    r1, _, r2, _ = v["rect"]
    assert all(r1 <= r <= r2 for r, _ in v["inputs"] + [v["cur"]])


def test_table_title_band_skips_rows_of_the_table_above(tmp_path):
    cells = {"S": {"A1": "코드", "B1": "금액", **{f"A{r}": f"c{r}" for r in range(2, 6)}, **{f"B{r}": 1 for r in range(2, 6)},
                   "A7": "종목", "B7": "수량", "A8": "x", "B8": 2, "D8": "=B8*2"}}
    book = make_book(tmp_path, cells, ("S", "T2", "A7:B8"))
    book.tables["T1"] = explain_cell.trace_cells.Table("S", (1, 1, 5, 2), ["코드", "금액"], 1, 0)
    blk = capture_views.block_of(book, "S", 8, 1)
    assert blk["rect"] == (7, 1, 8, 2)

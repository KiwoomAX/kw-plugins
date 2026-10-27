# explain_cell 시험. openpyxl 로 만든 작은 엑셀로 단계 · 도형 이름 · 출처 · 예외 입력을 본다.
#   uv run --no-project --with openpyxl --with pytest pytest tests/test_explain_cell.py
import sys
from pathlib import Path

import json
import os

import openpyxl
import pytest

ROOT = Path(__file__).resolve().parents[1] / "plugins/kw-dashboard/skills"
sys.path.insert(0, str(ROOT / "explaining-excel-cell/scripts"))
import explain_cell  # noqa: E402

# 표본 엑셀 경로는 사용자 PC 마다 달라 환경 변수로 받는다. 없으면 그 시험을 건너뛴다
DRAMATIC = Path(os.environ.get("EXPLAIN_DRAMATIC_XLSM", "/없음"))
FUNDINFO = Path(os.environ.get("EXPLAIN_FUNDINFO_XLSX", "/없음"))


def make(tmp_path, cells):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for sheet, values in cells.items():
        ws = wb.create_sheet(sheet)
        for coord, v in values.items():
            ws[coord] = v
    path = tmp_path / "b.xlsx"
    wb.save(path)
    return path


BOOK = {"원천": {"A1": "일자", "B1": "가격", "A2": 20250131, "B2": 100.0, "A3": 20250228, "B3": 110.0},
        "결과": {"A1": "수익률", "A2": "=(원천!B3-원천!B2)/원천!B2"}}


def test_steps_follow_sheet_reading_order_and_name_columns(tmp_path):
    data = explain_cell.build(make(tmp_path, BOOK), "결과", "A2")
    assert [s["addr"] for s in data["steps"]] == ["원천!B2", "원천!B3", "결과!A2"]
    assert data["steps"][2]["from"] == [0, 1]
    node = {n["id"]: n for n in data["nodes"]}
    assert node["원천!B"]["title"] == "가격" and node["원천!B"]["source"] == "수기 입력"
    assert node["결과!A"]["source"] == "계산"
    assert [g.get("at") for g in data["steps"][2]["segs"] if "v" in g] == ["원천!B3", "원천!B2", "원천!B2"]
    assert set(data) == {"title", "excel", "target", "nodes", "steps", "repeats", "views", "sources"}  # 화면이 쓰지 않는 항목은 없다
    assert all("stage" not in n for n in data["nodes"])


def test_cell_address_accepts_lowercase_and_dollars(tmp_path):
    path = make(tmp_path, BOOK)
    assert explain_cell.build(path, "결과", "$a$2")["steps"] == explain_cell.build(path, "결과", "A2")["steps"]


def test_missing_sheet_or_empty_cell_is_a_clear_error(tmp_path):
    path = make(tmp_path, BOOK)
    with pytest.raises(ValueError, match="시트 없음"):
        explain_cell.build(path, "없음", "A1")
    with pytest.raises(ValueError, match="빈 셀"):
        explain_cell.build(path, "결과", "Z9")


def test_long_running_balance_chain_does_not_hit_recursion_limit(tmp_path):
    cells = {"S": {"A1": 1.0, **{f"A{r}": f"=A{r - 1}+1" for r in range(2, 1801)}}}
    data = explain_cell.build(make(tmp_path, cells), "S", "A1800")
    assert len(data["steps"]) == 1800


def test_out_path_that_exists_is_refused(tmp_path):
    path, out = make(tmp_path, BOOK), tmp_path / "x.html"
    out.write_text("기존")
    with pytest.raises(SystemExit):
        explain_cell.main([str(path), "결과", "A2", "--out", str(out)])
    assert out.read_text() == "기존"


@pytest.mark.skipif(not DRAMATIC.exists(), reason="EXPLAIN_DRAMATIC_XLSM 이 가리키는 표본 엑셀이 없다")
def test_dramatic_starts_with_winus_import():
    data = explain_cell.build(DRAMATIC, "3004", "D26")
    steps = data["steps"]
    assert steps[0]["kind"] == "import" and steps[0]["to"] == "3004자료!G"
    assert [s["addr"] for s in steps[1:]] == ["3004자료!G175", "3004자료!G206", "3004자료!M17", "3004!D26"]
    assert [s["value"] for s in steps[1:]] == ["894.73", "921.3", "0.0296961", "0.0296961"]
    node = {n["id"]: n for n in data["nodes"]}
    # 출처 대응표를 읽으면(사내망) WINUS, 못 읽으면 DB 쿼리 연결이다. 표 이름은 공개 저장소라 적지 않는다
    source = node["3004자료!G"]["source"]
    assert source in ("WINUS", "DB 쿼리 연결") and node[steps[0]["node"]]["title"].startswith(f"{source} · ")


@pytest.mark.skipif(not FUNDINFO.exists(), reason="EXPLAIN_FUNDINFO_XLSX 가 가리키는 표본 엑셀이 없다")
def test_fundinfo_vlookup_follows_found_cell():
    data = explain_cell.build(FUNDINFO, "4.수치 수합", "C2")
    last = data["steps"][-1]
    assert last["value"] == "BNK든든한적격TDF2030(채혼-재간접)ClassC-P2"
    found = [g for g in last["segs"] if "v" in g and g["src"].upper().startswith("VLOOKUP")]
    assert found and found[0]["at"].startswith("연동!B")


@pytest.mark.skipif(not FUNDINFO.exists(), reason="EXPLAIN_FUNDINFO_XLSX 가 가리키는 표본 엑셀이 없다")
def test_fundinfo_pivot_reaches_its_sql_raw_row():
    steps = explain_cell.build(FUNDINFO, "4.수치 수합", "G2")["steps"]
    assert [s["addr"] for s in steps] == ["1.SQL_RAW!D399", "3. Pivot!C185", "4.수치 수합!G2"]
    assert steps[0]["value"] == steps[1]["value"] == "-0.00411359"


def test_names_file_overrides_shape_titles(tmp_path):
    data = explain_cell.build(make(tmp_path, BOOK), "결과", "A2", names={"원천!B": "단가"})
    node = {n["id"]: n for n in data["nodes"]}
    assert node["원천!B"]["title"] == "단가"
    assert data["steps"][0]["text"].startswith("단가 ")


def test_hints_give_rule_name_neighbours_block_label_and_formula(tmp_path):
    cells = {"S": {"A1": "월별성과", "B3": "펀드", "B5": 1.0, "B6": 2.0, "A6": 20250131, "C6": "=B6/B5-1"}}
    path = make(tmp_path, cells)
    book = explain_cell.Book(path)
    hints = {h["id"]: h for h in explain_cell.name_hints(book, explain_cell.build(book, "S", "C6"))}
    assert hints["S!B"]["rule"] == "펀드"
    assert hints["S!B"]["left_label"] == ["A1", "월별성과"]
    assert hints["S!B"]["above"] == [["B3", "펀드"]]
    assert hints["S!C"]["formula"] == "=B6/B5-1"


def test_hints_option_writes_file_without_html(tmp_path):
    path, out = make(tmp_path, BOOK), tmp_path / "hints.json"
    explain_cell.main([str(path), "결과", "A2", "--hints", str(out)])
    assert {h["id"] for h in json.loads(out.read_text(encoding="utf-8"))} == {"원천!B", "결과!A"}


def test_hints_with_missing_sheet_ends_with_one_line(tmp_path):
    with pytest.raises(SystemExit, match="시트 없음"):
        explain_cell.main([str(make(tmp_path, BOOK)), "없음", "A1", "--hints", str(tmp_path / "h.json")])


def test_missing_or_non_excel_file_ends_with_one_line(tmp_path):
    with pytest.raises(SystemExit, match="엑셀 파일 없음"):
        explain_cell.main([str(tmp_path / "없음.xlsx"), "S", "A1", "--out", str(tmp_path / "a.html")])
    bad = tmp_path / "b.xlsx:Zone.Identifier"  # 윈도우에서 받은 파일 옆에 생기는 표식 파일
    bad.write_text("[ZoneTransfer]")
    with pytest.raises(SystemExit, match="엑셀 열기 오류"):
        explain_cell.main([str(bad), "S", "A1", "--out", str(tmp_path / "b.html")])


def test_hints_path_that_exists_is_refused(tmp_path):
    out = tmp_path / "keep.json"
    out.write_text("기존")
    with pytest.raises(SystemExit):
        explain_cell.main([str(make(tmp_path, BOOK)), "결과", "A2", "--hints", str(out)])
    assert out.read_text() == "기존"


def test_cell_text_cannot_break_out_of_page_markup(tmp_path):
    book = {"S": {"A1": "x--><img src=x onerror=alert(1)>", "B1": "=A1"}}
    out = tmp_path / "p.html"
    explain_cell.main([str(make(tmp_path, book)), "S", "B1", "--out", str(out)])
    page = out.read_text(encoding="utf-8")
    assert "<img src=x" not in page
    comment = page.split("\n")[6]  # 템플릿 7행 주석이 데이터를 담지 않고 그 줄에서 닫힌다
    assert comment.startswith("<!-- explain_cell.py") and comment.endswith("-->") and "x--" not in comment


def test_view_has_play_bar_ids_notice_and_no_diagram_code(tmp_path):
    out = tmp_path / "p.html"
    explain_cell.main([str(make(tmp_path, BOOK)), "결과", "A2", "--out", str(out)])
    page = out.read_text(encoding="utf-8")
    assert 'id="speed" min="1" max="2.5" step="0.25" value="1.5"' in page  # explaining-code-flow 와 같은 속도 범위
    for button in ('id="play"', 'id="prev"', 'id="pp"', 'id="next"', 'id="first"', 'id="speedVal"', 'data-t="light"', 'data-t="dark"'):
        assert button in page
    assert 'id="again"' not in page  # 「이 단계 다시 보기」 단추는 두지 않는다
    assert "아래 보이는 화면은 엑셀을 그대로 분석하여 설명하는 화면입니다." in page
    assert "출처 확인" in page
    for gone in ("formulaFlow", "zoomTo", "moveTokens"):
        assert gone not in page


def test_view_data_json_has_no_raw_angle_bracket_and_keeps_cell_text(tmp_path):
    book = {"S": {"A1": "<img src=x onerror=alert(1)>", "B1": '=A1&"x"'}}
    out = tmp_path / "p.html"
    explain_cell.main([str(make(tmp_path, book)), "S", "B1", "--out", str(out)])
    page = out.read_text(encoding="utf-8")
    data = page.split('<script id="data" type="application/json">')[1].split("</script>")[0]
    assert "<" not in data and "\\u003cimg src=x" in data
    assert "<img src=x" not in page


def test_value_cell_in_formula_column_keeps_its_own_source(tmp_path):
    data = explain_cell.build(make(tmp_path, {"S": {"A1": 1.0, "A2": "=A1+1", "A3": "=A2+1"}}), "S", "A3")
    assert "손으로 넣은 값" in data["steps"][0]["text"] and data["steps"][0]["source"] == "수기 입력"


def test_external_file_reference_becomes_import_source(tmp_path):
    data = explain_cell.build(make(tmp_path, {"S": {"A1": 3.0, "B1": "=[1]Liquidity!C5*2", "C1": "=A1+B1"}}), "S", "C1")
    steps = data["steps"]
    imp = [s for s in steps if s["kind"] == "import"]
    assert len(imp) == 1 and imp[0]["to"] == "S!B"
    node = {n["id"]: n for n in data["nodes"]}
    assert node[imp[0]["node"]]["title"].startswith("외부 파일")
    b1 = next(s for s in steps if s["addr"] == "S!B1" and s["kind"] != "import")
    assert b1["kind"] == "value" and b1["source"].startswith("외부 파일") and not b1["from"]


def test_long_repeated_chain_is_folded(tmp_path):
    cells = {"S": {"A1": 1.0, **{f"A{r}": f"=A{r - 1}+1" for r in range(2, 21)}}}
    data = explain_cell.build(make(tmp_path, cells), "S", "A20")
    assert data["repeats"] == [{"first": 4, "last": 18, "count": 15, "sheet": "S", "col": 1, "rows": [5, 19], "kind": "chain"}]


def test_short_chain_is_not_folded(tmp_path):
    cells = {"S": {"A1": 1.0, **{f"A{r}": f"=A{r - 1}+1" for r in range(2, 7)}}}
    assert explain_cell.build(make(tmp_path, cells), "S", "A6")["repeats"] == []


RANGE_BOOK = {"S": {"A1": "=SUM(B3:B6)", "B3": "=C3*0", "B4": "=C4*2", "B5": "=C5*2", "B6": "=C6*2",
                    "C3": 1.0, "C4": 2.0, "C5": 3.0, "C6": 4.0}}


def test_range_without_saved_values_follows_first_cell_of_the_common_formula(tmp_path):
    # openpyxl 로 만든 파일은 저장값이 없다. 가장 많은 수식(B4:B6 의 =C*2)은 첫 칸 B4 가 대표이고,
    # 모양이 다른 B3(=C3*0)도 한 칸으로 따라간다
    data = explain_cell.build(make(tmp_path, RANGE_BOOK), "S", "A1")
    assert [s["addr"] for s in data["steps"]] == ["S!C3", "S!B3", "S!C4", "S!B4", "S!A1"]


def test_range_follows_first_nonzero_formula_cell_and_counts_the_rest(tmp_path):
    book = explain_cell.Book(make(tmp_path, RANGE_BOOK))
    book.values["S"].update({(1, 1): 18.0, (3, 2): 0.0, (4, 2): 4.0, (5, 2): 6.0, (6, 2): 8.0})
    data = explain_cell.build(book, "S", "A1")
    assert [s["addr"] for s in data["steps"]] == ["S!C3", "S!B3", "S!C4", "S!B4", "S!A1"]
    assert data["steps"][4]["from"] == [1, 3]
    note = data["steps"][4]["note"]
    assert "B4" in note and "모양이 다른 수식 1가지도 대표 한 칸씩 따라갔습니다" in note


def test_same_range_repeated_in_one_formula_is_explained_once(tmp_path):
    cells = {"S": {"A1": "=SUM(B3:B5)+SUM(B3:B5)", **{f"B{r}": f"=C{r}*2" for r in range(3, 6)}, "C3": 1.0, "C4": 2.0, "C5": 3.0}}
    note = explain_cell.build(make(tmp_path, cells), "S", "A1")["steps"][-1]["note"]
    assert note.count("나머지 2칸도 같은 계산입니다") == 1


def test_range_of_same_shaped_formulas_says_same_calculation(tmp_path):
    cells = {"S": {"A1": "=SUM(B3:B5)", **{f"B{r}": f"=C{r}*2" for r in range(3, 6)}, "C3": 1.0, "C4": 2.0, "C5": 3.0}}
    data = explain_cell.build(make(tmp_path, cells), "S", "A1")
    assert "나머지 2칸도 같은 계산입니다" in data["steps"][-1]["note"]
    rep = next(i for i, s in enumerate(data["steps"]) if s["addr"] == "S!B3")
    assert {"t": "B3:B5", "k": rep} in data["steps"][-1]["segs"]  # 화면이 대표 칸에서 범위 글자로 화살표를 그린다


def test_range_of_plain_values_becomes_one_range_step(tmp_path):
    cells = {"S": {"A1": "=SUM(B1:B3)", "B1": 1.0, "B2": 2.0, "B3": 3.0}}
    data = explain_cell.build(make(tmp_path, cells), "S", "A1")
    rng, last = data["steps"]
    assert rng["addr"] == "S!B1:B3" and rng["kind"] == "value" and rng["span"] == [1, 2, 3, 2]
    assert "값 칸 3개" in rng["text"] and "손으로 넣은 값" in rng["text"] and rng["source"] == "수기 입력"
    assert last["from"] == [0]


def test_range_without_any_cell_adds_no_step(tmp_path):
    data = explain_cell.build(make(tmp_path, {"S": {"A1": "=SUM(B1:B3)"}}), "S", "A1")
    assert [s["addr"] for s in data["steps"]] == ["S!A1"]


def test_range_with_different_shaped_formulas_follows_each_shape(tmp_path):
    cells = {"S": {"B1": "=C1*2", "B2": "=C2*2", "B3": "=C3*2", "B4": "=C4+1", "B5": "=C5+1",
                   **{f"C{r}": float(r) for r in range(1, 6)}, "A1": "=SUM(B1:B5)"}}
    addrs = {s["addr"] for s in explain_cell.build(make(tmp_path, cells), "S", "A1")["steps"]}
    assert {"S!B1", "S!B4"} <= addrs


def test_range_step_is_highlighted_as_a_span_in_views(tmp_path):
    cells = {"S": {"A1": "=SUM(B1:B3)", "B1": 1.0, "B2": 2.0, "B3": 3.0}}
    data = explain_cell.build(make(tmp_path, cells), "S", "A1")
    assert data["views"][0]["cur_span"] == [1, 2, 3, 2]
    assert data["views"][1]["input_spans"] == [[1, 2, 3, 2]] and data["views"][1]["inputs"] == [[1, 2]]


def test_whole_column_range_follows_representative(tmp_path):
    cells = {"S": {"A1": "=SUM(B:B)", "B1": "머리", "B2": 5.0, "B3": "=B2*2", "B4": "=B2*3"}}
    data = explain_cell.build(make(tmp_path, cells), "S", "A1")
    assert [s["addr"] for s in data["steps"]] == ["S!B2", "S!B3", "S!B4", "S!A1"]  # B3 · B4 는 모양이 달라 둘 다 따라간다


def test_range_representative_skips_header_and_parameter_formulas(tmp_path):
    # 열 전체 범위의 맨 위 수식은 날짜 머리나 매개변수 칸이 많다(Cash Management 설정해지 V6 · raw data C2).
    # 가장 많은 칸이 쓰는 수식의 첫 칸을 대표로 고르고, 모양이 다른 머리 수식은 그다음에 따라간다
    cells = {"S": {"A1": "=SUM(B:B)", "B1": "=TEXT(20261001,\"0\")", **{f"B{r}": f"=C{r}*2" for r in range(2, 6)},
                   **{f"C{r}": float(r) for r in range(2, 6)}}}
    data = explain_cell.build(make(tmp_path, cells), "S", "A1")
    assert [s["addr"] for s in data["steps"]] == ["S!B1", "S!C2", "S!B2", "S!A1"]
    assert data["steps"][-1]["from"] == [0, 2]


def test_value_range_of_a_query_table_shows_its_source(tmp_path):
    # SUMIF 가 더하는 범위가 WINUS 연결 표의 값 칸이면, 값 칸 위주여도 대표 한 칸을 따라가 출처를 보인다
    cells = {"S": {"A1": "=SUMIF(D:D,1,E:E)", "D1": "키", "E1": "금액", **{f"D{r}": 1.0 for r in range(2, 6)},
                   **{f"E{r}": float(r) for r in range(2, 6)}}}
    book = explain_cell.Book(make(tmp_path, cells))
    book.tables["표_Q"] = explain_cell.trace_cells.Table("S", (1, 4, 5, 5), ["키", "금액"], 1, 0, query=True)
    book.queries = {"표_Q": {"name": "q", "source": "WINUS", "tables": ["SCHEMA_A.DW_자산"]}}
    book.loaded_ranges = [{"sheet": "S", "area": (1, 4, 5, 5), "source": "WINUS", "kind": "ODBC 쿼리 테이블", "name": "q",
                           "tables": ["SCHEMA_A.DW_자산"]}]
    steps = explain_cell.build(book, "S", "A1")["steps"]
    assert [s["kind"] for s in steps][:2] == ["import", "value"] and steps[1]["source"] == "WINUS"
    assert "WINUS 연결 표의 값 칸" in steps[-1]["note"]


def query_book(tmp_path, cells, area):
    book = explain_cell.Book(make(tmp_path, cells))
    book.tables["표_Q"] = explain_cell.trace_cells.Table("S", area, ["펀드코드", "설정"], 1, 0, query=True)
    book.queries = {"표_Q": {"name": "q", "source": "WFS", "tables": ["T1@TESTLINK3"]}}
    book.loaded_ranges = [{"sheet": "S", "area": area, "source": "WFS", "kind": "ODBC 쿼리 테이블", "name": "q",
                           "tables": ["T1@TESTLINK3"]}]
    return book


def test_value_in_macro_result_range_shows_its_source(tmp_path):
    # 매크로(VBA)가 결과를 쓰는 열의 값은 수기 입력이 아니라 그 연결의 출처다
    book = explain_cell.Book(make(tmp_path, {"S": {"A1": "=B2*2", "B1": "금액", "B2": 5.0}}))
    book.loaded_ranges = [{"sheet": "S", "area": (2, 2, 2, 2), "source": "WINUS", "kind": "VBA ADODB SQL",
                           "name": "VBA M.P", "tables": []}]
    steps = explain_cell.build(book, "S", "A1")["steps"]
    assert next(s for s in steps if s["addr"] == "S!B2")["source"] == "WINUS"


def test_column_named_by_table_header_when_a_label_above_is_seen_first(tmp_path):
    # 설정해지 A열: 표 밖 A1 라벨을 먼저 보고 열 이름을 정해도, 표 안 데이터를 만나면 표 열 제목이 그 열의 이름이다
    cells = {"S": {"A1": "설정", "A3": "펀드코드", "A4": "2043", "B3": "설정", "B4": 5.0, "D1": "=A1&A4"}}
    steps = explain_cell.build(query_book(tmp_path, cells, (3, 1, 4, 2)), "S", "D1")["steps"]
    texts = [s["text"] for s in steps if s["addr"] in ("S!A1", "S!A4")]
    assert texts == ["펀드코드 값 설정 은(는) 손으로 넣은 값입니다", "펀드코드 값 2043 을(를) WFS 에서 가져옵니다"]


def test_table_header_cell_is_a_column_title_not_imported_data(tmp_path):
    # 수식이 표_Q[[#Headers],[설정]] 로 열 제목 글자를 읽으면, 그 칸은 WFS 데이터가 아니라 표 열 제목이다
    cells = {"S": {"A3": "펀드코드", "A4": "2043", "B3": "설정", "B4": 5.0, "D1": "=A4&B3"}}
    steps = explain_cell.build(query_book(tmp_path, cells, (3, 1, 4, 2)), "S", "D1")["steps"]
    head = next(s for s in steps if s["addr"] == "S!B3")
    assert head["source"] == "표 열 제목" and head["text"] == "설정 열 제목 글자 「설정」 입니다"


def test_range_of_mostly_values_skips_its_total_row(tmp_path):
    cells = {"S": {"A1": "=SUMIF(C:C,1,B:B)", **{f"B{r}": float(r) for r in range(2, 6)},
                   **{f"C{r}": 1.0 for r in range(2, 6)}, "B6": "=SUM(B2:B5)"}}
    data = explain_cell.build(make(tmp_path, cells), "S", "A1")
    assert [s["addr"] for s in data["steps"]][-1] == "S!A1"
    assert {s["addr"] for s in data["steps"]} == {"S!C2:C5", "S!B2:B5", "S!A1"}  # 합계 행(B6)은 값 칸 범위에 넣지 않는다


def make_pivot_book(tmp_path, source='ref="A1:C5" sheet="Src"', p_cells=None, table=False, extra=None):
    """원본 Src(코드 · 기간 · 값)를 행=코드, 열=기간, 값=합계로 묶은 피벗 P 를 엑셀 패키지에 끼워 넣는다."""
    import zipfile
    from openpyxl.worksheet.table import Table
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    cells = {"Src": {"A1": "코드", "B1": "기간", "C1": "값", "A2": "x", "B2": 1, "C2": 5.0, "A3": "x", "B3": 3, "C3": 7.0,
                     "A4": "y", "B4": 1, "C4": 2.0, "A5": "y", "B5": 1, "C5": 4.0},
             "P": {"A1": "행 레이블", "B1": 1, "C1": 3, "A2": "x", "B2": 5.0, "C2": 7.0, "A3": "y", "B3": 6.0,
                   "A4": "총합계", "B4": 11.0, **(p_cells or {})},
             **(extra or {})}
    for sheet, values in cells.items():
        ws = wb.create_sheet(sheet)
        for coord, v in values.items():
            ws[coord] = v
    if table:
        wb["Src"].add_table(Table(displayName="SrcTable", ref="A1:C5"))
    path = tmp_path / "p.xlsx"
    wb.save(path)
    m = 'xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"'
    pivot = (f'<pivotTableDefinition {m} name="p" cacheId="1"><location ref="A1:C4" firstHeaderRow="0" '
             'firstDataRow="1" firstDataCol="1"/><rowFields count="1"><field x="0"/></rowFields>'
             '<colFields count="1"><field x="1"/></colFields><dataFields count="1">'
             '<dataField name="합계 : 값" fld="2"/></dataFields></pivotTableDefinition>')
    cache = (f'<pivotCacheDefinition {m}><cacheSource type="worksheet"><worksheetSource {source}/>'
             '</cacheSource></pivotCacheDefinition>')
    rel = ('<?xml version="1.0"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/'
           'relationships"><Relationship Id="rId9" Type="{t}" Target="{target}"/></Relationships>')
    with zipfile.ZipFile(path, "a") as z:
        z.writestr("xl/pivotTables/pivotTable1.xml", pivot)
        z.writestr("xl/pivotTables/_rels/pivotTable1.xml.rels",
                   rel.format(t="pivotCacheDefinition", target="../pivotCache/pivotCacheDefinition1.xml"))
        z.writestr("xl/pivotCache/pivotCacheDefinition1.xml", cache)
        z.writestr("xl/worksheets/_rels/sheet2.xml.rels", rel.format(t="pivotTable", target="../pivotTables/pivotTable1.xml"))
    return path


def test_pivot_cell_follows_its_one_source_row(tmp_path):
    data = explain_cell.build(make_pivot_book(tmp_path), "P", "C2")
    assert [s["addr"] for s in data["steps"]] == ["Src!C3", "P!C2"]
    assert "Src!C3" in data["steps"][1]["note"]


def test_pivot_cell_over_several_rows_follows_a_representative(tmp_path):
    data = explain_cell.build(make_pivot_book(tmp_path), "P", "B3")
    assert [s["addr"] for s in data["steps"]] == ["Src!C4", "P!B3"]
    assert "2행" in data["steps"][1]["note"] and "나머지 1행" in data["steps"][1]["note"]


def test_pivot_grand_total_row_matches_every_row_of_its_column(tmp_path):
    data = explain_cell.build(make_pivot_book(tmp_path), "P", "B4")  # 총합계 행: 기간 1 인 원본 3행
    assert data["steps"][-1]["note"].count("3행") == 1


def test_pivot_over_whole_columns_follows_its_row(tmp_path):
    data = explain_cell.build(make_pivot_book(tmp_path, source='ref="A:C" sheet="Src"'), "P", "C2")
    assert [s["addr"] for s in data["steps"]] == ["Src!C3", "P!C2"]


def test_pivot_over_a_table_keeps_its_first_data_row(tmp_path):
    data = explain_cell.build(make_pivot_book(tmp_path, source='name="SrcTable"', table=True), "P", "B2")
    assert [s["addr"] for s in data["steps"]] == ["Src!C2", "P!B2"]


def test_pivot_rows_that_do_not_add_up_to_the_cell_are_not_followed(tmp_path):
    path = make_pivot_book(tmp_path, p_cells={"A3": "(비어 있음)", "B3": 7.0})  # 원본에 빈 코드 행이 없다
    assert [s["addr"] for s in explain_cell.build(path, "P", "B3")["steps"]] == ["P!B3"]


def test_range_over_pivot_cells_follows_a_pivot_cell(tmp_path):
    path = make_pivot_book(tmp_path, extra={"R": {"A1": "=SUM(P!B2:C3)"}})
    assert [s["addr"] for s in explain_cell.build(path, "R", "A1")["steps"]][-2:] == ["P!B2", "R!A1"]


def test_pivot_label_cell_is_not_followed(tmp_path):
    data = explain_cell.build(make_pivot_book(tmp_path), "P", "A2")
    assert [s["addr"] for s in data["steps"]] == ["P!A2"]


def test_circular_reference_through_a_range_is_noted_and_counted(tmp_path):
    cells = {"S": {"A1": "=B1+1", "B1": "=SUM(A1:A3)", "A2": "=C2*2", "A3": "=C3*2", "C2": 1.0, "C3": 2.0}}
    note = explain_cell.build(make(tmp_path, cells), "S", "A1")["steps"][-2]["note"]
    assert "순환 참조" in note and "수식 3칸" in note and "같은 계산" not in note


def test_circular_reference_stops_with_a_note(tmp_path):
    data = explain_cell.build(make(tmp_path, {"S": {"A1": "=B1+1", "B1": "=A1*2"}}), "S", "A1")
    assert [s["addr"] for s in data["steps"]] == ["S!B1", "S!A1"]
    assert "순환 참조" in data["steps"][0]["note"]


def test_dates_show_without_midnight_time(tmp_path):
    import datetime as dt
    assert explain_cell.show(dt.datetime(2016, 4, 4)) == "2016-04-04"
    assert explain_cell.show(dt.datetime(2016, 4, 4, 9, 30)) == "2016-04-04 09:30:00"
    data = explain_cell.build(make(tmp_path, {"S": {"A1": dt.datetime(2016, 4, 4), "B1": "=A1+1"}}), "S", "B1")
    assert [g["v"] for g in data["steps"][-1]["segs"] if "v" in g] == ["2016-04-04"]


def test_balance_chain_with_deposit_column_is_folded(tmp_path):
    cells = {"S": {"C2": 100.0, **{f"B{r}": float(r) for r in range(3, 21)},
                   **{f"C{r}": f"=C{r - 1}+B{r}" for r in range(3, 21)}}}
    data = explain_cell.build(make(tmp_path, cells), "S", "C20")
    (r,) = data["repeats"]
    shown = [s["addr"] for i, s in enumerate(data["steps"]) if not r["first"] <= i <= r["last"]]
    assert shown[:7] == ["S!C2", "S!B3", "S!C3", "S!B4", "S!C4", "S!B5", "S!C5"]
    assert shown[-2:] == ["S!B20", "S!C20"]
    assert r["count"] == 14 and r["rows"] == [6, 19]


def sum_of_inputs(tmp_path, n):
    """같은 열(수기 입력)의 값 칸 n 개를 한 칸이 더하는 엑셀."""
    cells = {"S": {**{f"A{r}": float(r) for r in range(1, n + 1)}, "B1": "=" + "+".join(f"A{r}" for r in range(1, n + 1))}}
    return explain_cell.build(make(tmp_path, cells), "S", "B1")


def test_six_consecutive_values_of_one_column_fold_after_the_first_three(tmp_path):
    data = sum_of_inputs(tmp_path, 7)
    assert [s["addr"] for s in data["steps"]][:3] == ["S!A1", "S!A2", "S!A3"]
    assert data["repeats"] == [{"first": 3, "last": 6, "count": 4, "sheet": "S", "col": 1, "rows": [4, 7],
                                "name": next(n["title"] for n in data["nodes"] if n["id"] == "S!A"),
                                "source": "수기 입력", "kind": "values"}]


def test_values_fold_starts_at_six_and_not_at_five(tmp_path):
    assert sum_of_inputs(tmp_path, 5)["repeats"] == []
    assert [r["count"] for r in sum_of_inputs(tmp_path, 6)["repeats"]] == [3]


def test_values_of_alternating_columns_do_not_fold(tmp_path):
    cells = {"S": {**{f"A{r}": float(r) for r in range(1, 5)}, **{f"B{r}": float(r) for r in range(1, 5)},
                   "C1": "=" + "+".join(f"A{r}+B{r}" for r in range(1, 5))}}
    assert explain_cell.build(make(tmp_path, cells), "S", "C1")["repeats"] == []


def test_sources_list_one_name_per_column_even_when_values_are_folded(tmp_path):
    data = sum_of_inputs(tmp_path, 7)
    manual = next(s for s in data["sources"] if s["kind"] == "수기 입력")
    assert [it["j"] for it in manual["items"]] == [0]  # 접힌 값도 같은 열이므로 이름은 하나이고 앞 3개 중 첫 단계로 간다
    views = data["views"]
    assert views[3]["cur"] == [7, 1] and views[3]["inputs"] == [[4, 1]]
    assert all(views[i] is None for i in (4, 5, 6))


@pytest.mark.skipif(not DRAMATIC.exists(), reason="EXPLAIN_DRAMATIC_XLSM 이 가리키는 표본 엑셀이 없다")
def test_dramatic_j60_folds_values_and_chain_into_two_items():
    data = explain_cell.build(DRAMATIC, "3004자료", "J60")
    assert sorted(r["kind"] for r in data["repeats"]) == ["chain", "values"]
    assert sum(v is not None for v in data["views"]) == 13


def test_value_range_across_blocks_makes_one_range_step_per_block(tmp_path):
    cells = {"S": {"A1": "=SUM(B1:B6)", "B1": 1.0, "B2": 2.0, "B3": 3.0, "B5": 5.0, "B6": 6.0}}  # 빈 행 4 가 두 덩어리로 가른다
    data = explain_cell.build(make(tmp_path, cells), "S", "A1")
    spans = [s["span"] for s in data["steps"] if s.get("span")]
    assert spans == [[1, 2, 3, 2], [5, 2, 6, 2]]
    assert data["steps"][-1]["from"] == [0, 1]


def test_all_mode_on_workbook_without_formulas_says_no_rules(tmp_path):
    with pytest.raises(SystemExit, match="계산 규칙 없음"):
        explain_cell.main([str(make(tmp_path, {"S": {"A1": 1.0}})), "--all", "--out", str(tmp_path / "x.html")])


def test_all_mode_refuses_existing_out_file(tmp_path):
    from conftest import THREE_SHEETS
    workbook = make(tmp_path, THREE_SHEETS)
    write_results(tmp_path, [{"sheet": "U", "excel_table": None, "row_order": "순위 순"}])
    (tmp_path / "out_r1.html").write_text("기존")
    with pytest.raises(SystemExit, match="출력 파일 중복"):
        explain_cell.main([str(workbook), "--all", "--results", str(tmp_path / "explain_results.json"),
                           "--out", str(tmp_path / "out.html")])
    assert (tmp_path / "out_r1.html").read_text() == "기존"


def test_all_with_hints_is_an_argument_error(tmp_path):
    with pytest.raises(SystemExit, match="인자 오류: --hints 는 셀 모드 전용"):
        explain_cell.main([str(make(tmp_path, {"S": {"C1": 3.0, "A1": "=C1*2"}})), "--all", "--hints", str(tmp_path / "h.json")])


def write_results(folder, results, excel="b.xlsx", names=None, raw=None, base=None, candidates_base=None):
    """결과 파일과, 2번째 실행이 대조하는 후보 파일(base_sheets 만)을 쓴다."""
    (folder / "explain_candidates.json").write_text(
        json.dumps({"excel": excel, "base_sheets": candidates_base if candidates_base is not None else (base or {})},
                   ensure_ascii=False), encoding="utf-8")
    path = folder / "explain_results.json"
    body = raw if raw is not None else json.dumps({"excel": excel, "base_sheets": base or {}, "results": results,
                                                   "names": names or {}}, ensure_ascii=False)
    path.write_bytes(body.encode("utf-8") if isinstance(body, str) else body)
    return path


def second_run(tmp_path, workbook, out=None):
    out = out or tmp_path / "out.html"
    explain_cell.main([str(workbook), "--all", "--results", str(out.parent / "explain_results.json"), "--out", str(out)])


def three(tmp_path):
    from conftest import THREE_SHEETS
    return make(tmp_path, THREE_SHEETS)


U = {"sheet": "U", "excel_table": None, "row_order": "순위 순"}


def test_first_run_writes_candidates_only(tmp_path, capsys):
    explain_cell.main([str(three(tmp_path)), "--all", "--out", str(tmp_path / "out.html")])
    assert "결과표 후보 시트 1개:" in capsys.readouterr().out
    cand = json.loads((tmp_path / "explain_candidates.json").read_text(encoding="utf-8"))
    assert cand["excel"] == "b.xlsx" and [c["sheet"] for c in cand["candidates"]] == ["U"]
    assert cand["candidates"][0]["shown"] is True and cand["candidates"][0]["tables"][0]["id"] == "U!A1:C4"
    assert cand["groups"][0] == {"id": "g1", "sheet": "S", "areas": ["A1:C4"], "header": ["코드", "이름", "값"],
                                 "title": "", "first_row": ["k2", "n2", "3"]}
    assert not list(tmp_path.glob("*.html"))


def test_second_run_writes_page_and_steps(tmp_path):
    workbook = three(tmp_path)
    write_results(tmp_path, [U])
    second_run(tmp_path, workbook)
    assert (tmp_path / "out_r1.html").exists() and any((tmp_path / "out_r1_steps").iterdir())


def test_two_results_make_two_pages(tmp_path):
    workbook = three(tmp_path)
    write_results(tmp_path, [U, {**U, "sheet": "T"}])
    second_run(tmp_path, workbook)
    assert (tmp_path / "out_r1.html").exists() and (tmp_path / "out_r2.html").exists()


def test_results_with_other_excel_name_fails(tmp_path):
    workbook = three(tmp_path)
    write_results(tmp_path, [U], excel="x.xlsx")
    with pytest.raises(SystemExit, match="결과표 파일의 엑셀 이름 다름: x.xlsx"):
        second_run(tmp_path, workbook)
    assert not list(tmp_path.glob("*.html"))


def test_results_with_missing_sheet_fails(tmp_path):
    workbook = three(tmp_path)
    write_results(tmp_path, [{**U, "sheet": "없음"}])
    with pytest.raises(SystemExit, match="결과표 시트 없음: 없음"):
        second_run(tmp_path, workbook)
    assert not list(tmp_path.glob("*.html"))


def test_results_with_bom_and_spaces_are_read(tmp_path):
    workbook = three(tmp_path)
    body = json.dumps({"excel": "b.xlsx", "results": [{**U, "sheet": " U "}], "names": {}}, ensure_ascii=False)
    write_results(tmp_path, None, raw=b"\xef\xbb\xbf" + body.encode("utf-8"))
    second_run(tmp_path, workbook)
    assert (tmp_path / "out_r1.html").exists()


def test_existing_page_fails_before_writing(tmp_path):
    workbook = three(tmp_path)
    write_results(tmp_path, [U, {**U, "sheet": "T"}])
    (tmp_path / "out_r2.html").write_text("기존")
    with pytest.raises(SystemExit, match="출력 파일 중복"):
        second_run(tmp_path, workbook)
    assert not (tmp_path / "out_r1.html").exists()


def test_failure_in_second_page_writes_nothing(tmp_path, monkeypatch):
    workbook = three(tmp_path)
    write_results(tmp_path, [U, {**U, "sheet": "T"}])
    real, calls = explain_cell.build_three_stage, []

    def flaky(*a, **k):
        calls.append(1)
        if len(calls) == 2:
            raise RuntimeError("두 번째 실패")
        return real(*a, **k)
    monkeypatch.setattr(explain_cell, "build_three_stage", flaky)
    with pytest.raises(RuntimeError):
        second_run(tmp_path, workbook)
    assert not (tmp_path / "out_r1.html").exists() and not (tmp_path / "out_r1_steps").exists()


def test_page_names_follow_out_name(tmp_path):
    workbook = three(tmp_path)
    folder = tmp_path / "a b"
    folder.mkdir()
    write_results(folder, [U])
    second_run(tmp_path, workbook, out=folder / "결과.html")
    assert (folder / "결과_r1.html").exists() and (folder / "결과_r1_steps").is_dir()


def test_many_candidate_sheets_mark_ten_shown(tmp_path):
    cells = {f"S{i:02d}": {"A1": "가", "B1": "나", "C1": "다", "A2": 1.5, "B2": "=A2*2", "C2": "=B2+1"} for i in range(21)}
    explain_cell.main([str(make(tmp_path, cells)), "--all", "--out", str(tmp_path / "out.html")])
    cand = json.loads((tmp_path / "explain_candidates.json").read_text(encoding="utf-8"))
    assert len(cand["candidates"]) == 21 and sum(c["shown"] for c in cand["candidates"]) == 10


DATED = {"2026년 1월1주": {"A1": "가", "B1": "나", "C1": "다", "A2": 1.5, "B2": "=A2*2", "C2": "=B2+1"},
         "2026년 1월2주": {"A1": "가", "B1": "나", "C1": "다", "A2": 2.5, "B2": "=A2*2", "C2": "=B2+1"}}


def test_base_argument_changes_analyzed_date_sheet(tmp_path):
    explain_cell.main([str(make(tmp_path, DATED)), "--all", "--base", "#년#월#주=2026년 1월1주", "--out", str(tmp_path / "o.html")])
    cand = json.loads((tmp_path / "explain_candidates.json").read_text(encoding="utf-8"))
    assert cand["base_sheets"] == {"#년#월#주": "2026년 1월1주"} and [c["sheet"] for c in cand["candidates"]] == ["2026년 1월1주"]


def test_base_argument_can_release_a_group(tmp_path):
    cells = {"TDF2030": DATED["2026년 1월1주"], "TDF2040": DATED["2026년 1월2주"]}
    explain_cell.main([str(make(tmp_path, cells)), "--all", "--base", "TDF#=없음", "--out", str(tmp_path / "o.html")])
    cand = json.loads((tmp_path / "explain_candidates.json").read_text(encoding="utf-8"))
    assert cand["base_sheets"] == {"TDF#": None} and cand["date_groups"] == []
    assert [c["sheet"] for c in cand["candidates"]] == ["TDF2030", "TDF2040"]


def test_unknown_base_is_an_argument_error(tmp_path):
    with pytest.raises(SystemExit, match="인자 오류: --base #년#월#주=2026년 9월1주"):
        explain_cell.main([str(make(tmp_path, DATED)), "--all", "--base", "#년#월#주=2026년 9월1주", "--out", str(tmp_path / "o.html")])


def test_results_base_must_match_candidates(tmp_path):
    workbook = make(tmp_path, DATED)
    write_results(tmp_path, [{**U, "sheet": "2026년 1월1주"}], base={"#년#월#주": "2026년 1월1주"}, candidates_base={})
    with pytest.raises(SystemExit, match="기준 시트가 후보 파일과 다름"):
        second_run(tmp_path, workbook)
    assert not list(tmp_path.glob("*.html"))


def test_results_base_with_spaces_matches_candidates(tmp_path):
    workbook = make(tmp_path, DATED)
    write_results(tmp_path, [{**U, "sheet": "2026년 1월1주"}], base={"#년#월#주": " 2026년 1월1주 "},
                  candidates_base={"#년#월#주": "2026년 1월1주"})
    second_run(tmp_path, workbook)
    assert "마지막 시트" not in (tmp_path / "out_r1.html").read_text(encoding="utf-8")


def test_skipped_date_sheet_as_result_fails(tmp_path):
    workbook = make(tmp_path, DATED)
    write_results(tmp_path, [{**U, "sheet": "2026년 1월1주"}])
    with pytest.raises(SystemExit, match="결과표 시트가 기준 시트가 아닌 날짜 시트: 2026년 1월1주"):
        second_run(tmp_path, workbook)


def test_excel_table_name_ignores_case_and_is_in_candidates(tmp_path):
    from conftest import make_book
    book = make_book(tmp_path, {"S": {"A1": "가", "B1": "나", "C1": "다", "A2": 1.5, "B2": "=A2*2", "C2": "=B2+1"}},
                     {"S": [("Table1", "A1:C2")]})
    name = Path(book.path).name
    explain_cell.main([book.path, "--all", "--out", str(tmp_path / "o.html")])
    cand = json.loads((tmp_path / "explain_candidates.json").read_text(encoding="utf-8"))
    assert cand["candidates"][0]["tables"][0]["excel_table"] == "Table1"
    write_results(tmp_path, [{"sheet": "S", "excel_table": "table1", "row_order": "순위 순"}], excel=name)
    second_run(tmp_path, book.path)
    assert (tmp_path / "out_r1.html").exists()
    write_results(tmp_path, [{"sheet": "S", "excel_table": "없는표", "row_order": "순위 순"}], excel=name)
    with pytest.raises(SystemExit, match="결과표 엑셀 표 없음: 없는표"):
        second_run(tmp_path, book.path, out=tmp_path / "o2.html")

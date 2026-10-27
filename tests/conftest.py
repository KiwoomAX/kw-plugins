# 설명기 3단계 시험들이 함께 쓰는 도우미. openpyxl 로 작은 엑셀을 만들어 Book 으로 읽는다.
import sys
from pathlib import Path

import openpyxl
from openpyxl.worksheet.table import Table

ROOT = Path(__file__).resolve().parents[1] / "plugins/kw-dashboard/skills"
sys.path.insert(0, str(ROOT / "explaining-excel-cell/scripts"))
import explain_cell  # noqa: E402

_count = [0]


THREE_SHEETS = {  # 원천 S → 계산 T(VLOOKUP) → 결과 U(두 배와 순위)
    "S": {"A1": "코드", "B1": "이름", "C1": "값", **{f"A{r}": f"k{r}" for r in (2, 3, 4)},
          **{f"B{r}": f"n{r}" for r in (2, 3, 4)}, **{f"C{r}": r * 1.5 for r in (2, 3, 4)}},
    "T": {"A1": "코드", "B1": "이름", "C1": "값", **{f"A{r}": f"k{r}" for r in (2, 3, 4)},
          **{f"B{r}": f"=VLOOKUP(A{r},S!$A$1:$C$4,2,FALSE)" for r in (2, 3, 4)},
          **{f"C{r}": f"=VLOOKUP(A{r},S!$A$1:$C$4,3,FALSE)" for r in (2, 3, 4)}},
    "U": {"A1": "코드", "B1": "두배", "C1": "순위", **{f"A{r}": f"k{r}" for r in (2, 3, 4)},
          **{f"B{r}": f"=T!C{r}*2" for r in (2, 3, 4)}, **{f"C{r}": f"=RANK(B{r},$B$2:$B$4)" for r in (2, 3, 4)}},
}


def all_tables(book):
    import tables
    return [t for s in book.order for t in tables.find_tables(book, s)]


def make_book(tmp_path, cells, excel_tables=None):
    """cells 는 {시트: {"A1": 값}}, excel_tables 는 {시트: [(이름, "A1:C4")]}. 같은 tmp_path 에 여러 번 만들어도 겹치지 않는다."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for sheet, values in cells.items():
        ws = wb.create_sheet(sheet)
        for coord, v in values.items():
            ws[coord] = v
        for name, ref in (excel_tables or {}).get(sheet, []):
            ws.add_table(Table(displayName=name, ref=ref))
    _count[0] += 1
    path = tmp_path / f"b{_count[0]}.xlsx"
    wb.save(path)
    return explain_cell.Book(str(path))


def three_sheets(tmp_path):
    return make_book(tmp_path, THREE_SHEETS)

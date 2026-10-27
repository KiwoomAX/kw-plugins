"""결과 폴더 이름을 짓고, .sql 과 index.xlsx 를 임시 폴더에서 검사한 뒤 결과 폴더로 옮긴다."""
import hashlib
import os
import pathlib
import re
import shutil
import tempfile
from dataclasses import dataclass

import openpyxl
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE

from _model import FILE_STATE_KINDS, Connection
from _secrets import _contains, find_leaks

HEADERS = ("원본 엑셀 파일명", "원본 경로", "항목 번호", "SQL 파일", "저장 위치", "연결 종류", "접속 대상",
           "출처 분류", "입력", "결과를 쓰는 곳", "비고")
INDEX = "index.xlsx"


@dataclass
class Result:
    win_path: str
    name: str
    conns: list[Connection]
    secrets: set[str]


def windows_path(p: str) -> str:
    """원본 경로를 Windows 절대 경로로 정규화한다. WSL 경로는 /mnt/<드라이브> 와 \\\\wsl.localhost 로 바꾼다."""
    if re.match(r"^[A-Za-z]:\\|^\\\\", p):
        return p
    p = os.path.abspath(p)
    m = re.match(r"^/mnt/([a-z])(/.*)?$", p)
    if m:
        return f"{m.group(1).upper()}:" + (m.group(2) or "\\").replace("/", "\\")
    if os.name == "nt":
        return p
    distro = os.environ.get("WSL_DISTRO_NAME", "Ubuntu")
    return f"\\\\wsl.localhost\\{distro}" + p.replace("/", "\\")


def _slug(text: str, sep: str) -> str:
    return sep.join(w.lower() for w in re.findall(r"[A-Za-z0-9]+", text))


def folder_name(win_path: str) -> str:
    stem = re.sub(r"\.[^.\\]+$", "", win_path.rsplit("\\", 1)[-1])
    digest = hashlib.sha1(win_path.lower().encode("utf-8")).hexdigest()[:6]
    slug = _slug(stem, "-")
    return f"{slug}-{digest}" if slug else digest


def sql_name(no: int, location: str) -> str:
    tail = _slug(location, "_")
    return f"q{no}_{tail}.sql" if tail else f"q{no}.sql"


def sql_text(source_name: str, c: Connection) -> str:
    clean = lambda s: s.replace(":", " ").replace("*/", "* /")
    head = [f"원본 {clean(source_name)}", f"저장 위치 {clean(c.location)}"]
    if c.inputs:
        head.append(f"입력 {clean(c.inputs)}")
    return "/*\n" + "\n".join(head) + "\n*/\n" + (c.sql or "") + "\n"


def is_locked(index: pathlib.Path) -> bool:
    if (index.parent / f"~${index.name}").exists():
        return True
    if not index.exists():
        return False
    try:
        with open(index, "r+b"):
            return False
    except PermissionError:
        return True


def existing_paths(index: pathlib.Path) -> set[str]:
    if not index.exists():
        return set()
    wb = openpyxl.load_workbook(index, read_only=True)
    try:  # read_only 통합문서는 닫아야 Windows 에서 뒤의 os.replace 가 파일을 바꿀 수 있다
        return {r[1] for r in wb.active.iter_rows(min_row=2, values_only=True) if r and r[1]}
    finally:
        wb.close()


def _row(r: Result, no: int, c: Connection, sql_rel: str) -> list:
    state = c.kind in FILE_STATE_KINDS
    row = [r.name, r.win_path, no, sql_rel, "" if state else c.location, c.kind, "" if state else c.target,
           "" if state else c.origin, c.inputs, c.output_range, ", ".join(c.notes)]
    # 셀에 넣을 수 없는 제어 문자 하나로 폴더 전체 실행이 멈추지 않게 지운다
    return [ILLEGAL_CHARACTERS_RE.sub("", v) if isinstance(v, str) else v for v in row]


def _source_rows(r: Result, tmp: pathlib.Path) -> list[list]:
    """원본 하나의 .sql 을 임시 폴더에 쓰고 행을 만든다."""
    folder = tmp / folder_name(r.win_path)
    rows = []
    for no, c in enumerate(r.conns, 1):
        rel = ""
        if c.sql:
            folder.mkdir(parents=True, exist_ok=True)
            name = sql_name(no, c.location)
            (folder / name).write_text(sql_text(r.name, c), encoding="utf-8")
            rel = f"{folder.name}/{name}"
        rows.append(_row(r, no, c, rel))
    return rows


def _leaked(r: Result, rows: list[list], tmp: pathlib.Path) -> bool:
    folder = tmp / folder_name(r.win_path)
    if folder.exists() and find_leaks(folder, r.secrets):
        return True
    return any(_contains(str(v), s) for row in rows for v in row if v is not None for s in r.secrets)


def _save_index(out: pathlib.Path, tmp: pathlib.Path, replace: set[str], new_rows: list[list]) -> pathlib.Path:
    index = out / INDEX
    if index.exists():
        wb = openpyxl.load_workbook(index)
        ws = wb.active
        for i in range(ws.max_row, 1, -1):
            if ws.cell(i, 2).value in replace or ws.cell(i, 2).value in {r[1] for r in new_rows}:
                ws.delete_rows(i)
    else:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "index"
        ws.append(list(HEADERS))
    for row in new_rows:
        ws.append(row)
        if row[3]:
            cell = ws.cell(ws.max_row, 4)
            cell.hyperlink = row[3]
            cell.style = "Hyperlink"
    path = tmp / INDEX
    wb.save(path)
    return path


def write(out: pathlib.Path, results: list[Result], replace: set[str]) -> list[str]:
    """결과를 임시 폴더에서 만들고 원본마다 누출을 검사한 뒤 out 으로 옮긴다. 누출된 원본 경로를 돌려준다."""
    out.mkdir(parents=True, exist_ok=True)
    tmp = pathlib.Path(tempfile.mkdtemp(prefix=f".{out.name}-", dir=out.parent))
    leaked = []
    try:
        new_rows = []
        for r in results:
            rows = _source_rows(r, tmp)
            if r.secrets and _leaked(r, rows, tmp):
                shutil.rmtree(tmp / folder_name(r.win_path), ignore_errors=True)
                err = Connection(kind="처리 오류", location="", notes=["비밀번호 누출 검사 실패"])
                rows = [_row(r, 1, err, "")]
                leaked.append(r.win_path)
            new_rows += rows
        index_tmp = _save_index(out, tmp, replace, new_rows)
        staged = out / f"{INDEX}.tmp"
        shutil.copyfile(index_tmp, staged)
        os.replace(staged, out / INDEX)
        for r in results:
            src = tmp / folder_name(r.win_path)
            dest = out / folder_name(r.win_path)
            if dest.exists():
                for old in dest.glob("*.sql"):
                    old.unlink()
            if not src.exists():
                continue
            dest.mkdir(exist_ok=True)
            for f in src.glob("*.sql"):
                shutil.move(str(f), dest / f.name)
    finally:
        shutil.rmtree(tmp, ignore_errors=True)
    return leaked


def _selfcheck() -> None:
    assert windows_path("/mnt/c/Users/a/b.xls") == r"C:\Users\a\b.xls"
    assert windows_path(r"C:\a\b.xls") == r"C:\a\b.xls"
    assert windows_path("/home/u/x").startswith("\\\\wsl")
    assert re.fullmatch(r"mmf-daily-[0-9a-f]{6}", folder_name(r"C:\x\샘플운용 MMF daily.xls"))
    # test_slug_hangul_only
    assert re.fullmatch(r"[0-9a-f]{6}", folder_name(r"C:\x\판매현황.xls"))
    # test_same_name_other_folder
    assert folder_name(r"C:\a\t.xlsx") != folder_name(r"C:\b\t.xlsx")
    assert sql_name(1, "VBA Module1.MMF현황") == "q1_vba_module1_mmf.sql"
    assert sql_name(2, "현황") == "q2.sql"
    # test_header_no_colon
    c = Connection(kind="VBA ADODB SQL", location="VBA M.P", sql="select 1", inputs="셀 U4 를 읽는 변수 s_ilja")
    head = sql_text("a:b.xls", c).split("*/")[0]
    assert head.startswith("/*") and ":" not in head
    with tempfile.TemporaryDirectory() as d:
        out = pathlib.Path(d, "excel-queries")
        r1 = Result(r"C:\x\t.xls", "t.xls", [Connection(kind="VBA ADODB SQL", location="VBA M.P", sql="select 1 from t", origin="WINUS")], set())
        assert write(out, [r1], set()) == []
        ws = openpyxl.load_workbook(out / "index.xlsx").active
        assert ws.cell(2, 4).hyperlink.target.endswith(".sql") and ws.cell(2, 8).value == "WINUS"
        assert existing_paths(out / "index.xlsx") == {r"C:\x\t.xls"}
        assert write(out, [r1], {r"C:\x\t.xls"}) == [] and openpyxl.load_workbook(out / "index.xlsx").active.max_row == 2
        leak = Result(r"C:\x\u.xls", "u.xls", [Connection(kind="VBA ADODB SQL", location="VBA M.P", sql="select 1", target="pw9secret")], {"pw9secret"})
        ok = Result(r"C:\x\v.xls", "v.xls", [Connection(kind="VBA ADODB SQL", location="VBA M.P", sql="select 2")], set())
        assert write(out, [leak, ok], set()) == [r"C:\x\u.xls"]
        rows = list(openpyxl.load_workbook(out / "index.xlsx").active.iter_rows(min_row=2, values_only=True))
        assert ("u.xls", "처리 오류") in [(r[0], r[5]) for r in rows] and not any("pw9secret" in str(v) for r in rows for v in r)
        # test_illegal_characters: 셀에 못 넣는 제어 문자가 있어도 실행이 멈추지 않는다
        bad = Result(r"C:\x\w.xls", "w.xls", [Connection(kind="DDE 링크", location="SupBook", target="app\x03topic")], set())
        assert write(out, [bad], set()) == []
        rows = list(openpyxl.load_workbook(out / "index.xlsx").active.iter_rows(min_row=2, values_only=True))
        assert ("w.xls", "apptopic") in [(r[0], r[6]) for r in rows]
        # test_is_locked
        size = (out / "index.xlsx").stat().st_size
        assert is_locked(out / "index.xlsx") is False and (out / "index.xlsx").stat().st_size == size
        (out / "~$index.xlsx").write_bytes(b"x")
        assert is_locked(out / "index.xlsx") is True


if __name__ == "__main__":
    _selfcheck()
    print("selfcheck ok")

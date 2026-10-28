"""엑셀 파일이나 폴더에서 외부 데이터 연결을 찾아 excel-queries/ 에 index.xlsx 와 .sql 로 남긴다."""
import pathlib
import sys
import tempfile

sys.path.insert(0, str(pathlib.Path(__file__).resolve().parent))

import argparse  # noqa: E402

from source_rules import load_rules  # noqa: E402

USAGE_ERROR = 4  # argparse 기본값 2 는 「index.xlsx 잠김」 과 겹쳐서 따로 둔다


class Parser(argparse.ArgumentParser):
    """인자 오류를 종료 코드 4 로 끝내는 파서. 두 실행 명령이 함께 쓴다."""

    def error(self, message):
        self.print_usage(sys.stderr)
        self.exit(USAGE_ERROR, f"{self.prog}: 인자 오류: {message}\n")


EXTENSIONS = {".xls", ".xlsx", ".xlsm", ".xlsb", ".xlt", ".xltx", ".xltm", ".xla", ".xlam"}
INSTALL = ('uv run --system-certs --no-project --with-requirements "<스킬>\\scripts\\requirements.txt" '
           'python "<스킬>\\scripts\\extract_queries.py" <경로>')
TEXT_FIELDS = ("location", "sql", "target", "inputs", "output_range")


def _files(paths: list[str], result_dir: pathlib.Path) -> list[pathlib.Path]:
    """입력 경로에서 엑셀 파일을 고른다. 결과 폴더 안의 파일(index.xlsx 등)은 뺀다."""
    skip = result_dir.resolve()
    out = []
    for p in map(pathlib.Path, paths):
        if p.is_dir():
            out += sorted(f for f in p.rglob("*") if f.is_file() and f.suffix.lower() in EXTENSIONS
                          and not f.name.startswith("~$") and not f.resolve().is_relative_to(skip))
        else:
            out.append(p)
    return out


def mask_fields(c, secrets: set[str]) -> None:
    """연결의 글자 칸에서 비밀값을 가린다. SQL 본문은 비밀번호 값으로 바꾸지 않는다.

    비밀번호와 같은 값이 SQL 에 있으면 바꾸는 순간 SQL 의미가 달라진다. 그런 SQL 은 누출 검사가
    그 원본을 「처리 오류」 행으로 돌린다.
    """
    from _secrets import mask
    for k in TEXT_FIELDS:
        v = getattr(c, k)
        if v:
            setattr(c, k, mask(v, set() if k == "sql" else secrets))


def dedupe(conns):
    """원본 안에서 공백을 하나로 줄인 SQL 이 같은 연결을 한 행으로 합친다."""
    out, seen = [], {}
    for c in conns:
        key = " ".join(c.sql.split()) if c.sql else None
        if key and key in seen:
            first = seen[key]
            if c.location not in first.location.split("; "):
                first.location += f"; {c.location}"
            first.notes = sorted(set(first.notes) | set(c.notes))
            continue
        if key:
            seen[key] = c
        out.append(c)
    return out


def _scan(path: pathlib.Path):
    """원본 하나를 읽어 (연결 목록, VBA 모듈 목록) 을 돌려준다. 실패는 파일 상태 행 하나로 바꾼다."""
    import olefile
    from _model import Connection
    from _ooxml import scan_zip
    from _xls import is_encrypted, scan_ole
    try:
        head = path.open("rb").read(8)
        if head.startswith(b"PK"):
            return scan_zip(path)
        if head.startswith(b"\xd0\xcf\x11\xe0"):
            with olefile.OleFileIO(str(path)) as ole:
                if is_encrypted(ole):
                    return [Connection(kind="열기 암호", location="")], []
            return scan_ole(path)
        return [Connection(kind="엑셀 형식 아님", location="")], []
    except Exception as e:  # 파일 하나의 예외는 행 하나로 남기고 다음 파일로 넘어간다
        return [Connection(kind="처리 오류", location="", notes=[type(e).__name__])], []


def _analyze(path: pathlib.Path, win_path: str, rules: dict):
    """원본 하나의 연결을 뽑아 비밀값을 가리고 합치고 분류한다. (연결 목록, 모은 비밀값)."""
    from _classify import classify
    from _secrets import collect
    conns, mods = _scan(path)
    texts = [m.source for m in mods] + [getattr(c, k) or "" for c in conns for k in ("target", "sql")]
    secrets = collect("\n".join(texts))
    for c in conns:
        mask_fields(c, secrets)
    conns = dedupe(conns)
    for c in conns:
        c.origin = classify(c, win_path, rules)
    return conns, secrets


def scan_file(path: pathlib.Path, rules: dict | None = None) -> list:
    """설명기가 호출하는 함수. 결과 파일을 쓰지 않고 가림 · 합치기 · 분류를 마친 연결 목록을 돌려준다.

    결과 파일 단계의 누출 검사를 메모리에서 한다. 모은 비밀값이 연결 칸에 남으면 그 원본은 「처리 오류」 한 행이다.
    rules 는 출처 대응표다. 호출한 쪽이 이미 읽었으면 넘겨서 공용 폴더를 다시 읽지 않게 한다.
    """
    from _index import windows_path
    from _model import Connection
    from _secrets import _contains
    conns, secrets = _analyze(path, windows_path(str(path)), rules if rules is not None else load_rules()[0])
    if any(_contains(getattr(c, k) or "", s) for c in conns for k in TEXT_FIELDS for s in secrets):
        return [Connection(kind="처리 오류", location="", notes=["비밀번호 누출 검사 실패"])]
    return conns


def main(argv: list[str]) -> int:
    try:
        import olefile  # noqa: F401
        import oletools  # noqa: F401
        import openpyxl  # noqa: F401
    except ImportError:
        print(f"olefile · oletools · openpyxl 이 필요하다. 이 형식으로 실행한다: {INSTALL}")
        return 1
    from _index import Result, existing_paths, is_locked, windows_path, write

    ap = Parser(prog="extract_queries.py")
    ap.add_argument("paths", nargs="+")
    ap.add_argument("--out", default=str(pathlib.Path.cwd() / "excel-queries"))
    ap.add_argument("--replace", action="store_true")
    args = ap.parse_args(argv)
    rules, state = load_rules()
    print(f"출처 대응표: {state}")
    out = pathlib.Path(args.out)
    if is_locked(out / "index.xlsx"):
        print(f"{out / 'index.xlsx'} 가 열려 있다. 닫고 다시 실행한다.")
        return 2
    files = _files(args.paths, out)
    win = {f: windows_path(str(f)) for f in files}
    already = existing_paths(out / "index.xlsx") & set(win.values())
    if already and not args.replace:
        for p in sorted(already):
            print(p)
        return 3

    results = []
    for f in files:
        conns, secrets = _analyze(f, win[f], rules)
        results.append(Result(win[f], f.name, conns, secrets))
    leaked = set(write(out, results, already))
    for r in results:
        unknown = sum(c.origin == "미판별" for c in r.conns)
        unsupported = sum("추출 미지원" in c.notes for c in r.conns)
        print(f"{r.name}\t{len(r.conns)}\t{unknown}\t{unsupported}\t{'누출 처리' if r.win_path in leaked else ''}")
    print(f"결과 폴더: {out}")
    return 0


def _selfcheck() -> None:
    import openpyxl
    from _model import Connection
    with tempfile.TemporaryDirectory() as d:
        d = pathlib.Path(d)
        src = d / "in"
        src.mkdir()
        (src / "fake.xls").write_text("<html></html>")
        (src / "~$lock.xlsx").write_bytes(b"PK")
        (src / "note.docx").write_bytes(b"PK")
        (src / "broken.xlsx").write_bytes(b"PK\x03\x04garbage")   # test_broken_file_row
        out = d / "excel-queries"
        assert main([str(src), "--out", str(out)]) == 0
        rows = list(openpyxl.load_workbook(out / "index.xlsx").active.iter_rows(min_row=2, values_only=True))
        assert sorted(r[5] for r in rows) == ["엑셀 형식 아님", "처리 오류"] and all(r[7] in (None, "") for r in rows)
        assert main([str(src), "--out", str(out)]) == 3
        assert main([str(src), "--out", str(out), "--replace"]) == 0
        (out / "~$index.xlsx").write_bytes(b"x")                   # test_locked_index
        before = (out / "index.xlsx").read_bytes()
        assert main([str(src), "--out", str(out), "--replace"]) == 2 and (out / "index.xlsx").read_bytes() == before
        (out / "~$index.xlsx").unlink()
        # test_skip_own_output: 결과 폴더가 입력 폴더 안이면 결과물을 다시 읽지 않는다
        (out / "x" / "sub").mkdir(parents=True)
        (out / "x" / "sub" / "copy.xlsm").write_text("<html></html>")
        assert main([str(d), "--out", str(out), "--replace"]) == 0
        rows = list(openpyxl.load_workbook(out / "index.xlsx").active.iter_rows(min_row=2, values_only=True))
        assert not any(r[0] in ("copy.xlsm", "index.xlsx") for r in rows)
        # test_sql_not_rewritten_by_secret: SQL 본문은 비밀번호 값으로 바꾸지 않는다(누출이면 처리 오류 행이 된다)
        c = Connection(kind="VBA ADODB SQL", location="VBA M.P", sql="select * from t where k = 1234", target="X 1234")
        mask_fields(c, {"1234"})
        assert c.sql == "select * from t where k = 1234" and c.target == "X ***"
        # test_usage_error_exit_code: 인자 오류는 「index.xlsx 잠김」(2) 과 다른 4 로 끝난다
        try:
            main(["--out"])
        except SystemExit as e:
            assert e.code == 4, e.code
        else:
            raise AssertionError("인자 오류인데 종료하지 않음")
        # test_dedupe_across_modules
        a = Connection(kind="VBA ADODB SQL", location="VBA A.P", sql="select 1 from t")
        b = Connection(kind="VBA ADODB SQL", location="VBA B.Q", sql="select  1 from t")
        [m] = dedupe([a, b])
        assert m.location == "VBA A.P; VBA B.Q"


if __name__ == "__main__":
    if sys.argv[1:] == ["--selfcheck"]:
        _selfcheck()
        print("selfcheck ok")
    else:
        sys.exit(main(sys.argv[1:]))

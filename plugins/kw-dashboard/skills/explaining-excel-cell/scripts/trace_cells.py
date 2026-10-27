"""엑셀의 셀 · 수식 · 표 · 피벗을 읽고, 결과 시트에서 참조를 거슬러 올라가 셀을 raw 값과 수식으로 나눈다.

설명기(explain_cell.py)가 엑셀 읽기 함수를 가져다 쓴다. 결과 시트의 모든 셀에서 시작해 수식이 참조하는 셀을 따라간다. 닿은 셀 중 값 셀은 raw,
수식 셀은 계산 코드로 옮길 대상이다. 닿지 않은 셀은 결과에 쓰이지 않는다.
"""
import posixpath
import re
import zipfile
import xml.etree.ElementTree as ET
from collections import deque
from dataclasses import dataclass, field

import openpyxl
from openpyxl.formula.tokenizer import Tokenizer
from openpyxl.utils import column_index_from_string

DYNAMIC_FUNCS = {"INDIRECT(", "OFFSET("}
_PART = re.compile(r"^\$?([A-Za-z]{1,3})?\$?(\d+)?$")


@dataclass
class Cell:
    kind: str  # "value", "formula", "pivot"(피벗 결과. 원본 범위에서 계산된다)
    formula: str | None = None
    reached: bool = False


@dataclass
class Table:
    sheet: str
    area: tuple  # (r1, c1, r2, c2) 머리 행과 합계 행 포함
    columns: list
    header_rows: int
    totals_rows: int
    query: bool = False  # 엑셀 데이터 연결(쿼리 테이블)이 채우는 표
    name: str = ""  # 엑셀에 적힌 표 이름(대소문자 그대로). 사전 키는 대문자다


@dataclass
class Pivot:
    sheet: str
    area: tuple
    source: str  # "시트!범위" 또는 표 · 정의된 이름
    rows: list = field(default_factory=list)  # 행 필드 번호(원본 범위의 몇 번째 열인가, 0 부터)
    cols: list = field(default_factory=list)  # 열 필드 번호
    data: list = field(default_factory=list)  # [(값 필드 번호, 묶는 방식 subtotal)]
    first: tuple = (0, 0, 0)  # location 의 firstHeaderRow · firstDataRow · firstDataCol


@dataclass
class Trace:
    cells: dict  # 시트 이름 -> {(행, 열): Cell}
    sheet_order: list
    results: list
    external_refs: dict = field(default_factory=dict)  # (시트, (행, 열)) -> [참조 문자열]
    dynamic_refs: set = field(default_factory=set)
    broken_refs: set = field(default_factory=set)
    precedents: dict = field(default_factory=dict)  # (시트, (행, 열)) -> [(시트, 영역)]
    ref_tokens: dict = field(default_factory=dict)  # (시트, (행, 열)) -> [(수식 속 참조 글자, [(시트, 영역)] 또는 None)]
    tables: dict = field(default_factory=dict)  # 표 이름 -> Table


def _formula_text(value):
    if isinstance(value, str) and value.startswith("="):
        return value
    text = getattr(value, "text", None)  # ArrayFormula
    if isinstance(text, str):
        return text if text.startswith("=") else "=" + text
    return None


def load_cells(path):
    wb = openpyxl.load_workbook(path, read_only=True)
    cells, order = {}, []
    for ws in wb.worksheets:
        order.append(ws.title)
        sheet = cells[ws.title] = {}
        for row in ws.iter_rows():
            for c in row:
                if c.value is None or not hasattr(c, "row"):
                    continue
                f = _formula_text(c.value)
                sheet[(c.row, c.column)] = Cell("formula", f) if f else Cell("value")
    names = {}
    for name, dn in wb.defined_names.items():
        names[name.upper()] = dn.attr_text
    wb.close()
    return cells, order, names


_NS = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
       "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
       "pr": "http://schemas.openxmlformats.org/package/2006/relationships"}


def _rels(z, part):
    """part 의 관계 파일을 읽어 {Id: (Type 끝 이름, 대상 경로)} 를 돌려준다."""
    rel_path = posixpath.join(posixpath.dirname(part), "_rels", posixpath.basename(part) + ".rels")
    if rel_path not in z.namelist():
        return {}
    out = {}
    for r in ET.fromstring(z.read(rel_path)).findall("pr:Relationship", _NS):
        target = r.get("Target")
        if target.startswith("/"):  # 패키지 루트 기준 절대 경로
            target = target[1:]
        else:
            target = posixpath.normpath(posixpath.join(posixpath.dirname(part), target))
        out[r.get("Id")] = (r.get("Type").rsplit("/", 1)[-1], target)
    return out


def load_package_parts(path):
    """셀 값으로는 보이지 않는 연결을 읽는다. 표(구조적 참조용)와 피벗(원본 범위)."""
    tables, pivots = {}, []
    with zipfile.ZipFile(path) as z:
        wb = ET.fromstring(z.read("xl/workbook.xml"))
        wb_rels = _rels(z, "xl/workbook.xml")
        for s in wb.find("m:sheets", _NS):
            sheet_part = wb_rels[s.get(f"{{{_NS['r']}}}id")][1]
            for kind, target in _rels(z, sheet_part).values():
                if kind == "table":
                    t = ET.fromstring(z.read(target))
                    cols = [c.get("name") for c in t.find("m:tableColumns", _NS)]
                    tables[t.get("displayName").upper()] = Table(
                        s.get("name"), _parse_area(t.get("ref")), cols,
                        int(t.get("headerRowCount", "1")), int(t.get("totalsRowCount", "0")),
                        t.get("tableType") == "queryTable", t.get("displayName"))
                elif kind == "pivotTable":
                    p = ET.fromstring(z.read(target))
                    caches = [tg for k, tg in _rels(z, target).values() if k == "pivotCacheDefinition"]
                    if not caches:
                        continue
                    src = ET.fromstring(z.read(caches[0])).find("m:cacheSource/m:worksheetSource", _NS)
                    if src is None:  # 외부 데이터 원본 피벗
                        continue
                    source = src.get("name") or f"'{src.get('sheet')}'!{src.get('ref')}"
                    loc = p.find("m:location", _NS)
                    fields = lambda tag: [int(f.get("x")) for f in p.findall(f"m:{tag}/m:field", _NS)]
                    pivots.append(Pivot(
                        s.get("name"), _parse_area(loc.get("ref")), source,
                        [x for x in fields("rowFields") if x >= 0], [x for x in fields("colFields") if x >= 0],
                        [(int(d.get("fld")), d.get("subtotal", "sum")) for d in p.findall("m:dataFields/m:dataField", _NS)],
                        tuple(int(loc.get(k, "0")) for k in ("firstHeaderRow", "firstDataRow", "firstDataCol"))))
    return tables, pivots


_SPECIAL = {"#ALL", "#DATA", "#HEADERS", "#TOTALS", "#THIS ROW"}


def _structured_area(ref, sheet, pos, tables):
    """'T1[[#This Row],[값]]' · 'T1[값]' · '[@값]' 같은 구조적 참조를 (시트, 영역) 으로 바꾼다."""
    m = re.match(r"^([^\[\]]*)\[(.*)\]$", ref)
    if not m:
        return None
    name, inner = m.group(1), m.group(2)
    if name:
        table = tables.get(name.upper())
    else:
        table = next((t for t in tables.values() if t.sheet == sheet and _inside(pos, t.area)), None)
    if table is None:
        return None
    this_row = inner.startswith("@")
    parts = re.findall(r"\[([^\[\]]*)\]", inner) or [inner.lstrip("@")]
    if this_row and inner.startswith("@["):
        parts = re.findall(r"\[([^\[\]]*)\]", inner[1:])
    specials = {p.upper() for p in parts if p.upper() in _SPECIAL}
    names = [p for p in parts if p.upper() not in _SPECIAL]
    r1, c1, r2, c2 = table.area
    data_r1, data_r2 = r1 + table.header_rows, r2 - table.totals_rows
    if this_row or "#THIS ROW" in specials:
        rows = (pos[0], pos[0])
    elif "#ALL" in specials:
        rows = (r1, r2)
    elif "#HEADERS" in specials:
        rows = (r1, r1)
    elif "#TOTALS" in specials:
        rows = (r2, r2)
    else:
        rows = (data_r1, data_r2)
    if names:
        try:
            idx = [table.columns.index(n.replace("'", "")) for n in names]
        except ValueError:
            return None
        cols = (c1 + min(idx), c1 + max(idx))
    else:
        cols = (c1, c2)
    return table.sheet, (rows[0], cols[0], rows[1], cols[1])


def _inside(pos, area):
    r1, c1, r2, c2 = area
    return r1 <= pos[0] <= r2 and c1 <= pos[1] <= c2


def _parse_area(text):
    """'A1' · '$A$1:$B$5' · 'C:D' · '3:5' 를 (r1, c1, r2, c2) 로 바꾼다. 열린 끝은 None."""
    ends = []
    for part in text.split(":"):
        m = _PART.match(part)
        if not m or (m.group(1) is None and m.group(2) is None):
            return None
        col = column_index_from_string(m.group(1).upper()) if m.group(1) else None
        r = int(m.group(2)) if m.group(2) else None
        ends.append((r, col))
    if len(ends) == 1:
        ends.append(ends[0])
    if len(ends) != 2:
        return None
    (r1, c1), (r2, c2) = ends
    return r1, c1, r2, c2


def _split_ref(ref, current_sheet):
    """'시트!범위' 를 (시트, 범위) 로 나눈다. 시트가 없으면 현재 시트."""
    if "!" not in ref:
        return current_sheet, ref
    sheet, area = ref.rsplit("!", 1)
    if sheet.startswith("'") and sheet.endswith("'"):
        sheet = sheet[1:-1].replace("''", "'")
    return sheet, area


def make_resolver(cells, tables, names):
    """참조 문자열 하나를 [(시트, 영역)] 으로 바꾸는 함수를 만든다. 외부면 "external", 못 풀면 "broken"."""
    def resolve(ref, sheet, pos):
        if "#REF" in ref.upper():
            return "broken"
        if "!" not in ref and "[" in ref:
            found = _structured_area(ref, sheet, pos, tables)
            return [found] if found else "broken"
        target, area_text = _split_ref(ref, sheet)
        if target.startswith("[") or "]" in target:
            return "external"
        if target not in cells:
            return "broken"
        area = _parse_area(area_text)
        if area is not None:
            return [(target, area)]
        if "!" in ref:
            return "broken"
        if area_text.upper() in tables:
            table = tables[area_text.upper()]
            r1, c1, r2, c2 = table.area
            return [(table.sheet, (r1 + table.header_rows, c1, r2 - table.totals_rows, c2))]
        named = names.get(area_text.upper())
        if not named:
            return "broken"
        out = []
        for part in named.split(","):
            n_sheet, n_area = _split_ref(part.strip(), sheet)
            parsed = _parse_area(n_area)
            if n_sheet in cells and parsed:
                out.append((n_sheet, parsed))
        return out or "broken"
    return resolve


def trace(path, results):
    cells, order, names = load_cells(path)
    missing = [r for r in results if r not in cells]
    if missing:
        raise ValueError(f"엑셀에 없는 결과 시트: {', '.join(missing)}")
    t = Trace(cells=cells, sheet_order=order, results=list(results))
    tables, pivots = load_package_parts(path)
    t.tables = tables
    pivot_source = {}
    for pv in pivots:
        if pv.sheet in cells and pv.area:
            for pos, cell in cells[pv.sheet].items():
                if _inside(pos, pv.area):
                    cell.kind, cell.formula = "pivot", f"피벗 원본: {pv.source}"
                    pivot_source[(pv.sheet, pos)] = pv.source
    cache = {}

    resolve = make_resolver(cells, tables, names)

    def expand(sheet, area):
        key = (sheet, area)
        if key not in cache:
            r1, c1, r2, c2 = area
            cache[key] = [
                (r, c) for (r, c) in cells[sheet]
                if (r1 is None or min(r1, r2) <= r <= max(r1, r2))
                and (c1 is None or min(c1, c2) <= c <= max(c1, c2))
            ]
        return cache[key]

    queue = deque()

    def reach(sheet, pos):
        cell = cells[sheet][pos]
        if not cell.reached:
            cell.reached = True
            if cell.kind != "value":
                queue.append((sheet, pos))

    for sheet in results:
        for pos in list(cells[sheet]):
            reach(sheet, pos)

    while queue:
        sheet, pos = queue.popleft()
        key = (sheet, pos)
        if key in pivot_source:
            refs = [pivot_source[key]]
        else:
            refs = _references(cells[sheet][pos].formula, t, key)
        for item in refs:
            ref, narrow = (item[0], item[1:]) if isinstance(item, tuple) else (item, None)
            found = resolve(ref, sheet, pos)
            if narrow and isinstance(found, list):
                found = [(s, a) for s, area in found for a in _narrow(area, *narrow)]
            t.ref_tokens.setdefault(key, []).append((ref, found if isinstance(found, list) else None))
            if found == "external":
                t.external_refs.setdefault(key, []).append(ref)
            elif found == "broken":
                t.broken_refs.add(key)
            else:
                for target, area in found:
                    t.precedents.setdefault(key, []).append((target, area))
                    for p in expand(target, area):
                        reach(target, p)
    return t


def formula_cells(t, sheets):
    """시트들에서 결과에 쓰이는 계산 셀(수식 · 피벗)마다 (시트, 행, 열, 수식, 참조 범위 목록)을 돌려준다.

    참조 범위는 (시트, r1, c1, r2, c2) 이고 열린 끝은 None 이다. 외부 파일 참조는 시트 자리에 "외부:<참조>" 를 둔다.
    """
    out = []
    for sheet in sheets:
        for (r, c), cell in sorted(t.cells.get(sheet, {}).items()):
            if not cell.reached or cell.kind == "value":
                continue
            key = (sheet, (r, c))
            refs = [(s, *area) for s, area in t.precedents.get(key, [])]
            refs += [(f"외부:{ref}", None, None, None, None) for ref in t.external_refs.get(key, [])]
            out.append((sheet, r, c, cell.formula, refs))
    return out


LOOKUP_AXIS = {"VLOOKUP(": "col", "HLOOKUP(": "row"}  # 표 범위 중 키와 돌려줄 열(행)만 읽는 함수


def _narrow(area, axis, index):
    """VLOOKUP 표 범위를 키 열과 index 번째 열로 좁힌다(HLOOKUP 은 행)."""
    r1, c1, r2, c2 = area
    if axis == "col" and c1 is not None:
        start = min(c1, c2)
        return [(r1, start, r2, start), (r1, start + index - 1, r2, start + index - 1)]
    if axis == "row" and r1 is not None:
        start = min(r1, r2)
        return [(start, c1, start, c2), (start + index - 1, c1, start + index - 1, c2)]
    return [area]


def _references(formula, t, key):
    """수식이 읽는 참조 문자열 목록. VLOOKUP · HLOOKUP 의 표 범위는 (참조, 축, 번호) 로 돌려준다."""
    try:
        tokens = Tokenizer(formula).items
    except Exception:
        t.broken_refs.add(key)
        return []
    refs, stack = [], []  # stack: 열린 함수마다 [이름, 인수 번호, 인수별 토큰 목록]
    for tok in tokens:
        if tok.type == "FUNC" and tok.subtype == "OPEN":
            if tok.value.upper() in DYNAMIC_FUNCS:
                t.dynamic_refs.add(key)
            stack.append([tok.value.upper(), 0, [[]]])
            continue
        if tok.type == "FUNC" and tok.subtype == "CLOSE":
            name, _, args = stack.pop() if stack else ("", 0, [])
            if name in LOOKUP_AXIS and len(args) >= 3:
                table, idx = args[1], args[2]
                if len(table) == 1 and table[0][0] == "RANGE" and len(idx) == 1 and idx[0][0] == "NUMBER" \
                        and isinstance(refs[table[0][1]], str):
                    refs[table[0][1]] = (refs[table[0][1]], LOOKUP_AXIS[name], int(float(idx[0][2])))
            if stack:
                stack[-1][2][-1].append(("FUNC", None, None))
            continue
        if tok.type == "SEP" and tok.subtype == "ARG" and stack:
            stack[-1][1] += 1
            stack[-1][2].append([])
            continue
        if tok.type == "OPERAND" and tok.subtype in ("RANGE", "ERROR"):
            if tok.subtype == "ERROR" and "#REF" not in tok.value.upper():
                continue
            refs.append(tok.value)
            if stack:
                stack[-1][2][-1].append(("RANGE", len(refs) - 1, tok.value))
        elif stack and tok.type != "WHITE-SPACE":
            stack[-1][2][-1].append((tok.subtype or tok.type, None, tok.value))
    return refs

"""엑셀 셀 하나의 값이 어떻게 계산됐는지 사람의 시선 순서로 재생하는 HTML 화면을 만든다.

    python explain_cell.py <엑셀> <시트> <셀> --out <html> [--title <제목>]

엑셀 파일만 읽는다. 값은 엑셀에 저장된 계산 결과이고, 수식은 다시 계산하지 않는다.
결과 셀의 수식에서 값이 들어가는 참조를 거슬러 올라가며 단계를 쌓고, 열 하나를 도형 하나로 묶는다.
"""
import argparse
import datetime as dt
import json
import re
import sys
import zipfile
from collections import Counter
from pathlib import Path

import openpyxl
from openpyxl.utils.exceptions import InvalidFileException

HERE = Path(__file__).resolve().parent
import trace_cells  # noqa: E402
from formula_segments import jsonable, segments  # noqa: E402
from openpyxl.formula.tokenizer import Tokenizer  # noqa: E402
from openpyxl.formula.translate import Translator  # noqa: E402
from query_sources import file_name, load_connections, load_external_files, query_tables  # noqa: E402

sys.setrecursionlimit(20000)  # 위 행을 계속 참조하는 열(드라매틱 엑셀은 1,722 행)도 끝까지 따라간다


def col_name(n):
    s = ""
    while n:
        n, m = divmod(n - 1, 26)
        s = chr(65 + m) + s
    return s


def a1(r, c):
    return f"{col_name(c)}{r}"


def parse_a1(text):
    m = re.fullmatch(r"\$?([A-Za-z]{1,3})\$?(\d+)", text.strip())
    if not m:
        raise ValueError(f"셀 주소 오류: {text}")
    return int(m.group(2)), openpyxl.utils.column_index_from_string(m.group(1).upper())


def show(v):
    if isinstance(v, dt.datetime):  # 엑셀 날짜는 자정 시각이 붙어 오므로 시각이 0 이면 날짜만 보인다
        return v.date().isoformat() if v.time() == dt.time() else v.isoformat(sep=" ")
    v = jsonable(v)
    if isinstance(v, float):
        return str(int(v)) if v.is_integer() else f"{v:.6g}"
    return "" if v is None else str(v)


class Book:
    """엑셀 하나에서 수식 · 저장값 · 표 · 피벗 · 쿼리 연결을 읽어 둔다."""

    def __init__(self, path):
        self.path = path
        self.cells, self.order, names = trace_cells.load_cells(path)
        self.tables, pivots = trace_cells.load_package_parts(path)
        self.resolve = trace_cells.make_resolver(self.cells, self.tables, names)
        self.pivot_at = {(pv.sheet, pos): pv for pv in pivots if pv.sheet in self.cells and pv.area
                         for pos in self.cells[pv.sheet] if trace_cells._inside(pos, pv.area)}
        wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
        self.values = {ws.title: {(c.row, c.column): c.value for row in ws.iter_rows() for c in row
                                  if c.value is not None and hasattr(c, "row")} for ws in wb.worksheets}
        wb.close()
        # 쿼리 · 매크로가 결과를 쓰는 범위와 그 출처. 시트를 정할 수 없는 매크로 연결은 unplaced 에 둔다
        self.loaded_ranges, self.unplaced, self.source_rules_state = load_connections(path, self.values)
        self.queries = query_tables(self.tables, self.loaded_ranges)
        self.external_files = load_external_files(path)

    def loaded_source(self, sheet, pos):
        """셀이 쿼리 · 매크로 결과 범위 안이면 그 출처, 아니면 None."""
        r, c = pos
        return next((x["source"] for x in self.loaded_ranges if x["sheet"] == sheet
                     and x["area"][0] <= r <= x["area"][2] and x["area"][1] <= c <= x["area"][3]), None)

    def sheet_values(self, name):
        return self.values.get(name, {})

    def formula(self, sheet, pos):
        if (sheet, pos) in self.pivot_at:
            return f"피벗 원본: {self.pivot_at[(sheet, pos)].source}"
        cell = self.cells[sheet].get(pos)
        return cell.formula if cell is not None and cell.kind == "formula" else None

    def external_file(self, formula, sheet, pos):
        """수식이 다른 엑셀 파일을 참조하면 그 파일 이름을 돌려준다(`[1]` 번호는 외부 링크의 파일 이름으로 바꾼다)."""
        try:
            tokens = Tokenizer(formula).items
        except Exception:
            return None
        for tok in tokens:
            if tok.type == "OPERAND" and tok.subtype == "RANGE" and self.resolve(tok.value, sheet, pos) == "external":
                m = re.search(r"\[([^\]]+)\]", tok.value)
                key = m.group(1) if m else ""
                return self.external_files.get(key) or (file_name(key) if key and not key.isdigit() else f"외부 링크 {key}")
        return None

    def table_at(self, sheet, pos):
        return next(((name, t) for name, t in self.tables.items()
                     if t.sheet == sheet and trace_cells._inside(pos, t.area)), (None, None))


def build(path, sheet, cell, title="", names=None):
    import capture_views  # capture_views 가 이 모듈의 a1 · show 를 쓰므로 호출할 때 가져온다
    book = path if isinstance(path, Book) else Book(path)
    names = names or {}
    if sheet not in book.cells:
        raise ValueError(f"시트 없음: {sheet}")
    r0, c0 = parse_a1(cell)
    if (r0, c0) not in book.cells[sheet]:
        raise ValueError(f"빈 셀: {sheet}!{a1(r0, c0)}")
    target = (sheet, r0, c0)
    steps, nodes, step_of, imported, visiting, by_col = [], {}, {}, {}, set(), {}
    range_step_of, extras = {}, []  # range_step_of: (시트, 값 칸 범위) -> 범위 단계 번호. extras: 대표와 모양이 다른 수식 칸
    shapes, s_now = {}, [sheet]  # shapes: 수식의 상대 위치 글자 캐시. s_now: 지금 대표 칸을 고르는 수식의 시트

    retitled = {}  # 도형 id -> 처음 붙인 이름. 같은 열의 표 안 칸을 나중에 만나 표 열 제목으로 바꾼 도형

    def header_cell(s, pos):
        """엑셀 표 맨 위의 열 제목 칸인가. 수식이 [#Headers] 로 이 글자를 읽기도 한다(설정해지 E6)."""
        name, t = book.table_at(s, pos)
        return t is not None and pos[0] < t.area[0] + t.header_rows

    def node_title(s, r, c):
        """표 안이면 표의 열 머리, 아니면 같은 열에서 위로 올라가며 처음 만나는 글자 셀."""
        name, t = book.table_at(s, (r, c))
        if t is not None:
            return t.columns[c - t.area[1]]
        for rr in range(r - 1, 0, -1):
            v = book.sheet_values(s).get((rr, c))
            if isinstance(v, str):
                return v
        return f"{s} {col_name(c)}열"

    def source_of(s, pos, calc, ext=None):
        if ext:
            return f"외부 파일({ext})"
        if (s, pos) in book.pivot_at:
            return "피벗"
        if calc:
            return "계산"
        if header_cell(s, pos):
            return "표 열 제목"
        return book.loaded_source(s, pos) or "수기 입력"

    def add_import(s, r, c, nid):
        """쿼리 테이블에서 온 원천이면, 그 표의 첫 원천 단계 앞에 가져오기 단계를 하나 둔다."""
        name, t = book.table_at(s, (r, c))
        if t is None or not t.query or name not in book.queries or name in imported:
            return
        q = book.queries[name]
        iid = f"src:{name}"
        nodes[iid] = {"id": iid, "sheet": s, "col": "", "kind": "import", "source": q["source"],
                      "title": f"{q['source']} · {', '.join(q['tables']) or q['name']}",
                      "range": f"엑셀 연결 「{q['name']}」"}
        tr1, tc1, tr2, tc2 = t.area
        where = f" 의 {', '.join(q['tables'])} 에서" if q["tables"] else " 에서"
        imported[name] = len(steps)
        steps.append({"node": iid, "to": nid, "from": [], "sheet": s, "row": tr1, "col": c,
                      "addr": f"{s}!{a1(tr1, c)}", "value": "", "kind": "import", "segs": None,
                      "text": f"{s} 시트의 표 {a1(tr1, tc1)}:{a1(tr2, tc2)} 은(는) {q['source']}{where} "
                              f"엑셀 연결 「{q['name']}」로 가져옵니다",
                      "note": "연결 문자열은 보이지 않습니다. 출처와 읽는 표 이름만 보입니다."})

    def add_external_import(s, r, c, nid, ext):
        """다른 엑셀 파일을 참조하는 셀은, 그 파일의 첫 원천 단계 앞에 「외부 파일」 가져오기 단계를 하나 둔다."""
        iid = f"ext:{ext}"
        if iid in nodes:
            return
        nodes[iid] = {"id": iid, "sheet": s, "col": "", "kind": "import", "source": f"외부 파일({ext})",
                      "title": f"외부 파일 · {ext}", "range": "다시 계산하려면 이 파일을 함께 첨부해야 합니다"}
        steps.append({"node": iid, "to": nid, "from": [], "sheet": s, "row": r, "col": c,
                      "addr": f"{s}!{a1(r, c)}", "value": "", "kind": "import", "segs": None,
                      "text": f"{s} 시트의 {a1(r, c)} 값은 외부 파일 「{ext}」에서 가져옵니다",
                      "note": "이 값은 엑셀이 다른 파일에서 읽어 온 값입니다. 다시 계산하려면 이 파일을 함께 첨부해야 합니다."})

    def pattern(rs, p):
        """수식을 상대 위치 기준으로 맞춘 글자. 같은 열의 행마다 같은 계산이면 같은 글자가 된다."""
        if (rs, p) not in shapes:
            f = book.formula(rs, p) or ""
            try:
                shapes[(rs, p)] = Translator(f, origin=a1(*p)).translate_formula("ZZ100000")
            except Exception:
                shapes[(rs, p)] = f
        return shapes[(rs, p)]

    def query_value(rs, inside, text):
        """값 칸 범위가 쿼리 연결 표의 값이면, 첫 값 칸을 대표로 따라가 출처(가져오기 단계)를 보인다."""
        vals = book.sheet_values(rs)
        for p in inside:
            name, t = book.table_at(rs, p)
            if (t is not None and t.query and name in book.queries and p[0] > t.area[0]
                    and vals.get(p) is not None and not isinstance(vals.get(p), str)):
                at = a1(*p) if rs == s_now[0] else f"{rs}!{a1(*p)}"
                return (rs, *p, f"범위 {text} 는 {book.queries[name]['source']} 연결 표의 값 칸이라 대표로 {at} 한 칸만 "
                                f"따라갔습니다. 나머지 칸은 따라가지 않았습니다.")
        return None

    def cells_in(rs, area):
        """범위 안에 있는 엑셀 칸 위치. 행마다 범위를 가진 긴 사슬도 빠르도록, 시트마다 열 -> 행 목록 색인을 한 번만 만든다."""
        r1, c1, r2, c2 = area
        if rs not in by_col:
            by_col[rs] = {}
            for rr, cc in sorted(book.cells.get(rs, {})):
                by_col[rs].setdefault(cc, []).append(rr)
        return [(rr, cc) for cc in (range(c1, c2 + 1) if c1 is not None else by_col[rs]) for rr in by_col[rs].get(cc, [])
                if r1 is None or r1 <= rr <= r2]

    def representative(s, ref, text):
        """범위 안 수식 칸 중 가장 많은 칸이 쓰는 수식의 칸에서, 값이 0도 빈 값도 아닌 첫 칸을 대표로 고른다.
        (시트, 행, 열, 설명) 이거나 None.

        열 전체 범위의 맨 위 수식은 날짜 머리나 매개변수 칸인 일이 많아(Cash Management 설정해지 V6 ·
        raw data C2), 첫 수식 칸을 고르면 실제로 더한 행 대신 머리를 따라가게 된다.
        값 칸이 수식 칸보다 많으면 따라가지 않는다. 그런 범위의 수식 칸은 대개 맨 아래 합계 행이라,
        대표로 보이면 실제로 더한 값 칸 대신 합계를 원천처럼 보이게 된다(에코자원화 L:L, 값 178칸 · 합계 행 하나).
        다만 그 값 칸이 쿼리 연결 표의 값이면 대표 한 칸을 따라가 출처를 보인다(Cash Management SUMIF 의 raw data U:U).
        """
        s_now[0] = s
        rs = ref["sheet"]
        cells, vals = book.cells.get(rs, {}), book.sheet_values(rs)
        inside = cells_in(rs, ref["area"])
        calc = lambda p: cells[p].kind == "formula" or (rs, p) in book.pivot_at  # 피벗 칸도 계산된 칸이다
        found = sorted(p for p in inside if calc(p))
        plain = sum(1 for p in inside if not calc(p) and not isinstance(vals.get(p), str))
        if not found or plain >= len(found):
            rep = query_value(rs, sorted(inside), text)
            return rep if rep and (rs, *rep[1:3]) not in visiting else None
        loop = [p for p in found if (rs, *p) in visiting]  # 지금 따라가는 셀로 되돌아가는 순환은 고르지 않는다
        cand = [p for p in found if p not in loop]
        cyc = f" {', '.join(a1(*p) for p in loop)} 은(는) 이 칸을 다시 참조하는 순환 참조라 따라가지 않았습니다." if loop else ""
        if not cand:
            return (None, None, None, f"범위 {text} 의 수식 칸이 모두 순환 참조라 여기서 추적을 멈춥니다.")
        shapes_count = Counter(pattern(rs, p) for p in cand)
        common = shapes_count.most_common(1)[0][0]  # 같은 수가 여럿이면 위쪽 수식
        for shape, _ in [x for x in shapes_count.most_common() if x[0] != common][:MAX_OTHER_SHAPES]:  # 모양이 다른 수식도 한 칸씩 따라간다
            group = [p for p in cand if pattern(rs, p) == shape]
            extras.append((rs, *next((p for p in group if vals.get(p) not in (None, 0, "")), group[0])))
        cand = [p for p in cand if pattern(rs, p) == common]
        rep = next((p for p in cand if vals.get(p) not in (None, 0, "")), cand[0])
        at = a1(*rep) if rs == s else f"{rs}!{a1(*rep)}"
        rest = len(found) - 1
        if not rest:
            return (rs, *rep, f"범위 {text} 의 수식 칸 {at} 을(를) 따라갔습니다.")
        try:
            f0 = book.formula(rs, rep)
            same = not loop and all(Translator(f0, origin=a1(*rep)).translate_formula(a1(*p)) == book.formula(rs, p)
                                    for p in found)
        except Exception:
            same = False
        tail = (f"나머지 {rest}칸도 같은 계산입니다." if same else
                f"모양이 다른 수식 {len(extras)}가지도 대표 한 칸씩 따라갔습니다. 나머지 칸은 따라가지 않았습니다." if extras else
                f"나머지 {rest}칸은 따라가지 않았습니다.")
        return (rs, *rep, f"범위 {text} 의 수식 {len(found)}칸 중 대표로 {at} 한 칸만 따라갔습니다. {tail}{cyc}")

    def range_steps(rs, ref, text):
        """값 칸 위주 범위를 칸마다 펴지 않고, 캡처 덩어리마다 단계 하나로 만든다. 단계 번호 목록(값 칸이 없으면 빈 목록)."""
        vals, cells = book.sheet_values(rs), book.cells.get(rs, {})
        plain = sorted(p for p in cells_in(rs, ref["area"]) if cells[p].kind != "formula" and (rs, p) not in book.pivot_at
                       and vals.get(p) is not None and not header_cell(rs, p))
        groups = {}
        for p in plain:  # 범위가 빈 행이나 열을 건너 여러 덩어리에 걸치면 덩어리마다 한 단계가 된다
            groups.setdefault(capture_views.block_of(book, rs, *p)["rect"], []).append(p)
        out = []
        for part in groups.values():
            box = (min(r for r, _ in part), min(c for _, c in part), max(r for r, _ in part), max(c for _, c in part))
            if (rs, box) not in range_step_of:
                r0, c0 = part[0]
                nid = f"{rs}!{col_name(c0)}"
                if nid not in nodes:
                    nodes[nid] = {"id": nid, "sheet": rs, "col": col_name(c0), "kind": "value",
                                  "title": names.get(nid) or node_title(rs, r0, c0), "range": f"{rs} 시트 {col_name(c0)}열",
                                  "source": source_of(rs, (r0, c0), False), "_table": book.table_at(rs, (r0, c0))[1] is not None}
                add_import(rs, r0, c0, nid)
                src, first = source_of(rs, (r0, c0), False), show(vals[(r0, c0)])
                where = f"{a1(box[0], box[1])}:{a1(box[2], box[3])}" if len(part) > 1 else a1(box[0], box[1])
                how = "손으로 넣은 값입니다" if src == "수기 입력" else f"{src} 에서 가져옵니다"
                steps.append({"node": nid, "from": [], "sheet": rs, "row": r0, "col": c0, "addr": f"{rs}!{where}", "value": first,
                              "text": f"{nodes[nid]['title']} 범위 {where} 의 값 칸 {len(part)}개는 {how}", "kind": "value",
                              "segs": None, "source": src, "span": list(box),
                              "note": f"범위 {text} 의 값 칸 {len(part)}개(첫 칸 {a1(r0, c0)} = {first})를 한꺼번에 읽습니다. 칸마다 단계를 만들지 않았습니다."})
                range_step_of[(rs, box)] = len(steps) - 1
            out.append(range_step_of[(rs, box)])
        return out

    def pivot_source_row(s, r, c):
        """피벗 값 칸이 원본의 어느 행들을 묶었는지 찾아 대표 한 칸을 고른다. (시트, 행, 열, 설명) 이거나 None.

        행 · 열 필드가 하나씩까지이고 값 필드가 하나인 피벗만 푼다. 그 칸의 행 레이블(위치의 첫 열)과
        열 레이블(머리 행)이 원본 행의 그 필드 값과 같은 행을 모은다. 총합계처럼 원본에 없는 레이블은 조건에서 뺀다.
        모은 행을 피벗의 묶는 방식으로 다시 계산해 피벗 저장값과 다르면 따라가지 않는다. 「(비어 있음)」 레이블이나
        읽지 못한 피벗 모양에서 틀린 행을 원천처럼 보이지 않기 위해서다.
        """
        pv = book.pivot_at[(s, (r, c))]
        hr, dr, dc = pv.first
        top, left = pv.area[0], pv.area[1]
        if len(pv.data) != 1 or len(pv.rows) > 1 or len(pv.cols) > 1 or r < top + dr or c < left + dc:
            return None
        found = book.resolve(pv.source, s, (r, c))
        if not isinstance(found, list) or len(found) != 1:
            return None
        ss, (sr1, sc1, sr2, sc2) = found[0]
        if sc1 is None:
            return None
        sv, vals = book.sheet_values(ss), book.sheet_values(s)
        key = lambda v: str(v.date() if isinstance(v, dt.datetime) and v.time() == dt.time() else
                            int(v) if isinstance(v, float) and v.is_integer() else v).upper()
        used = sorted({rr for rr, cc in sv if sc1 <= cc <= sc2})
        if not used:
            return None
        if pv.source.upper() in book.tables:  # 표 이름 원본은 머리 행을 뺀 데이터 영역으로 풀린다
            first = sr1
        else:  # 시트 범위 원본은 첫 행이 머리 행이다. 열 전체(A:C)면 그 열에서 처음 쓰인 행이다
            first = (sr1 if sr1 is not None else used[0]) + 1
        body = [rr for rr in used if rr >= first and (sr2 is None or rr <= sr2)]
        conds = []
        for f, label in [(f, vals.get((r, left))) for f in pv.rows] + [(f, vals.get((top + hr, c))) for f in pv.cols]:
            if label is not None and key(label) in {key(sv[(rr, sc1 + f)]) for rr in body if (rr, sc1 + f) in sv}:
                conds.append((sc1 + f, key(label)))
        dcol = sc1 + pv.data[0][0]
        rows = [rr for rr in body if (rr, dcol) in sv and all(key(sv.get((rr, col))) == k for col, k in conds)]
        nums = [sv[(rr, dcol)] for rr in rows if isinstance(sv[(rr, dcol)], (int, float)) and not isinstance(sv[(rr, dcol)], bool)]
        calc = {"sum": lambda: sum(nums), "count": lambda: len(rows), "average": lambda: sum(nums) / len(nums),
                "max": lambda: max(nums), "min": lambda: min(nums)}
        how = {"sum": "더했습니다", "count": "개수를 셌습니다", "average": "평균을 냈습니다",
               "max": "최댓값을 골랐습니다", "min": "최솟값을 골랐습니다"}
        want = vals.get((r, c))
        if (not rows or pv.data[0][1] not in calc or (not nums and pv.data[0][1] != "count")
                or not isinstance(want, (int, float)) or abs(calc[pv.data[0][1]]() - want) > 1e-9 * max(1, abs(want))):
            return None
        at = a1(rows[0], dcol) if ss == s else f"{ss}!{a1(rows[0], dcol)}"
        if len(rows) == 1:
            return (ss, rows[0], dcol, f"피벗이 원본 {at} 한 행의 값을 그대로 옮긴 칸입니다.")
        how = how[pv.data[0][1]]
        return (ss, rows[0], dcol, f"피벗이 원본 {len(rows)}행을 {how}. 대표로 {at} 한 행만 따라갔고, 나머지 {len(rows) - 1}행은 따라가지 않았습니다.")

    def walk(s, r, c):
        key = (s, r, c)
        if key in step_of:
            return step_of[key]
        visiting.add(key)
        formula = book.formula(s, (r, c))
        ext = book.external_file(formula, s, (r, c)) if formula else None
        if ext:
            formula = None  # 다른 파일에서 읽어 온 값은 계산이 아니라 원천으로 다룬다
        nid = f"{s}!{col_name(c)}"
        in_table = book.table_at(s, (r, c))[1] is not None
        if nid not in nodes:
            nodes[nid] = {"id": nid, "sheet": s, "col": col_name(c), "kind": "calc" if formula else "value",
                          "title": names.get(nid) or (title if key == target and title else node_title(s, r, c)),
                          "range": f"{s} 시트 {col_name(c)}열",
                          "source": source_of(s, (r, c), bool(formula), ext), "_table": in_table}
        elif in_table and not nodes[nid]["_table"] and nid not in names and not (title and nid == f"{sheet}!{col_name(c0)}"):
            # 같은 열에 표 밖 라벨(설정해지 A4)과 표 안 데이터(A7 펀드코드)가 있으면 표 열 제목이 그 열의 이름이다
            retitled.setdefault(nid, nodes[nid]["title"])
            nodes[nid].update(title=node_title(s, r, c), _table=True)
            if not header_cell(s, (r, c)):
                nodes[nid]["source"] = source_of(s, (r, c), bool(formula), ext)
        raw = segments(formula, s, (r, c), book.sheet_values, book.resolve)["segments"] if formula else []
        targets = [(g["ref"]["sheet"], g["ref"]["area"][0], g["ref"]["area"][1]) for g in raw
                   if "v" in g and g["ref"]["area"][0] is not None]
        notes, rep_of, range_of = [], {}, {}  # rep_of: 범위 조각 번호 -> 대표 칸. range_of: 범위 조각 번호 -> 값 범위 단계 번호 목록
        if formula and formula.startswith("피벗 원본:"):  # 피벗은 원본 범위 대신 그 칸이 묶은 원본 행을 따라간다
            rep = pivot_source_row(s, r, c)
            if rep and rep[:3] not in visiting:
                targets.append(rep[:3])
                notes.append(rep[3])
                rep_of[next(i for i, g in enumerate(raw) if g.get("ref"))] = rep[:3]
        elif formula:
            for i, g in enumerate(raw):
                if "t" in g and g.get("ref"):
                    extras.clear()
                    rep = representative(s, g["ref"], g["t"])
                    if rep:
                        notes.append(rep[3])
                        if rep[0] is not None:
                            targets.append(rep[:3])
                            rep_of[i] = rep[:3]
                            targets.extend(extras)
                    else:
                        ks = range_steps(g["ref"]["sheet"], g["ref"], g["t"])
                        if ks:
                            range_of[i] = ks
        for t in [t for t in targets if t in visiting]:  # 지금 따라가는 셀로 되돌아오면 끝없이 돌므로 멈춘다
            notes.append(f"{t[0]}!{a1(t[1], t[2])} 을(를) 다시 참조하는 순환 참조라 여기서 추적을 멈춥니다.")
        targets = [t for t in targets if t not in visiting]
        parents = []
        for ts, tr, tc in sorted(targets, key=lambda x: (x[1], x[2])):  # 시트를 읽는 순서
            p = walk(ts, tr, tc)
            if p not in parents:
                parents.append(p)
        parents.extend(k for ks in range_of.values() for k in ks if k not in parents)
        value = book.sheet_values(s).get((r, c))
        if formula:
            segs = []
            for i, g in enumerate(raw):
                if i in rep_of:
                    segs.append({"t": g["t"], "k": step_of[rep_of[i]]})
                elif i in range_of:
                    segs.append({"t": g["t"], "k": range_of[i][0]})
                elif "v" in g:
                    at = a1(g["ref"]["area"][0], g["ref"]["area"][1])
                    at = at if g["ref"]["sheet"] == s else f"{g['ref']['sheet']}!{at}"
                    k = step_of.get((g["ref"]["sheet"], g["ref"]["area"][0], g["ref"]["area"][1]))
                    v = book.sheet_values(g["ref"]["sheet"]).get((g["ref"]["area"][0], g["ref"]["area"][1]))
                    segs.append({"at": at, "src": g["src"], "v": show(v), "k": k})  # 날짜를 원본 값으로 보이려고 다시 읽는다
                else:
                    segs.append({"t": g["t"]})
            if segs and segs[0].get("t") == "=":
                segs = segs[1:]
            expr = "".join(f"[{g['v']}]" if "v" in g else g["t"] for g in segs)
            text = f"{nodes[nid]['title']}: {expr} = {show(value)}"
        else:
            segs = None
            if ext:
                add_external_import(s, r, c, nid, ext)
            elif not header_cell(s, (r, c)):  # 열 제목 글자는 연결이 가져온 데이터가 아니다
                add_import(s, r, c, nid)
            src = source_of(s, (r, c), False, ext)  # 한 열에 수식 셀과 값 셀이 섞일 수 있어 셀마다 정한다
            text = (f"{nodes[nid]['title']} 값 {show(value)} 은(는) 손으로 넣은 값입니다" if src == "수기 입력"
                    else f"{nodes[nid]['title']} 열 제목 글자 「{show(value)}」 입니다" if src == "표 열 제목"
                    else f"{nodes[nid]['title']} 값 {show(value)} 을(를) {src} 에서 가져옵니다")
        step_of[key] = len(steps)
        steps.append({"node": nid, "from": parents, "sheet": s, "row": r, "col": c, "addr": f"{s}!{a1(r, c)}",
                      "value": show(value), "text": text, "kind": "calc" if formula else "value", "segs": segs,
                      "source": source_of(s, (r, c), bool(formula), ext)})
        if notes:
            steps[-1]["note"] = " ".join(dict.fromkeys(notes))  # 같은 범위가 수식에 여러 번 나와도 안내는 한 번만 적는다
        visiting.discard(key)
        return step_of[key]

    walk(*target)
    for st in steps:  # 이름을 바꾼 도형은, 바꾸기 전에 쓴 단계 문장 앞머리도 새 이름으로 맞춘다
        old = retitled.get(st["node"])
        if old and st["text"].startswith(old):
            st["text"] = nodes[st["node"]]["title"] + st["text"][len(old):]
    for n in nodes.values():
        n.pop("_table", None)
    last = steps[-1]
    last["text"] = (f"결과 셀 「{title}」의 {last['addr']} 칸에 {last['value']} 이(가) 들어갑니다" if title
                    else f"결과 셀 {last['addr']} 칸에 {last['value']} 이(가) 들어갑니다")

    repeats = find_repeats(steps, list(nodes.values()))
    return {"title": title, "excel": Path(book.path).name, "target": f"{sheet}!{a1(r0, c0)}",
            "nodes": list(nodes.values()), "steps": steps,
            "repeats": repeats, "views": capture_views.build_views(book, steps, repeats),
            "sources": capture_views.build_sources(list(nodes.values()), steps)}


def render_page(data):
    """화면 템플릿에 데이터 JSON 을 넣어 HTML 글자를 돌려준다. 셀 글자가 HTML 을 끊지 못하게 < 를 JSON 이스케이프로 바꾼다."""
    page = (HERE.parent / "assets/explain_view.html").read_text(encoding="utf-8")
    return page.replace("__DATA__", json.dumps(data, ensure_ascii=False).replace("<", "\\u003c"))


PREVIEW_ROWS = 13  # 미리보기 행 수(목업에서 쓴 값, 근거 없음)
LINES_MAX = 6  # 카드마다 보이는 계산 설명 줄 수(목업에서 쓴 값, 근거 없음)
FALLBACK_END = " 셀은 수식으로 계산됩니다."  # plain_text.FALLBACK 의 끝
GUIDE = {"text": "분석 결과가 실제와 다르면, 알고 계신 엑셀 구조를 아래 예시처럼 간단히 설명해 주세요. "
                 "클로드가 설명을 참고해 다시 분석합니다.",
         "example": "예시: 「A」 시트의 펀드별 수익률 표는 WINUS 에서 엑셀로 추출해 붙여 넣는 값입니다. "
                    "「B」 시트는 그 값으로 펀드별 비중을 계산하고, 펀드 행마다 정리해 순위를 매깁니다. "
                    "「C」 시트의 지난주 열은 지난주 시트의 값을 매주 손으로 붙여 넣습니다."}


def analyze(book, base_override=None):
    """엑셀 전체의 표와 표 사이 읽기 관계. 날짜 시트 군의 기준 시트가 아닌 날짜 시트는 표를 찾지 않는다."""
    import tables as tb
    sheets, groups = tb.analyzed_sheets(book, base_override)
    skipped = [s for s in book.order if s not in sheets]
    found = [t for s in sheets for t in tb.find_tables(book, s)]
    reads, odd_refs = tb.read_map(book, found, set(skipped))
    return {"sheets": sheets, "skipped": skipped, "date_groups": groups, "tables": found,
            "groups": tb.group_tables(found), "reads": reads, "odd_refs": odd_refs}


def _area_text(area):
    return f"{a1(area[0], area[1])}:{a1(area[2], area[3])}"


def _zone(book, t):
    """「Passive (A~L열)」. 구역 이름은 표의 열 범위에서 시트 1행의 첫 글자 셀이다."""
    vals = book.sheet_values(t["sheet"])
    _, c1, _, c2 = t["area"]
    name = next((vals[(1, c)].strip() for c in range(c1, c2 + 1) if isinstance(vals.get((1, c)), str) and vals[(1, c)].strip()), "")
    cols = f"{col_name(c1)}~{col_name(c2)}열" if c1 != c2 else f"{col_name(c1)}열"
    return f"{name} ({cols})" if name else cols


def _preview(book, ts):
    t = max(ts, key=lambda x: x["area"][2] - x["area"][0])  # 행이 가장 많은 표. 같으면 앞 표
    r1, c1, r2, c2 = t["area"]
    vals = book.sheet_values(t["sheet"])
    rows = [[{"v": show(vals.get((r, c))), "formula": bool(book.formula(t["sheet"], (r, c)))} for c in range(c1, c2 + 1)]
            for r in range(r1, min(r2, r1 + PREVIEW_ROWS - 1) + 1)]
    note = (f"같은 모양 표 {len(ts)}개 중 이 표({_area_text(t['area'])})를 예로 보여 드립니다. 이 표의 행이 가장 많습니다."
            if len(ts) > 1 else None)
    return {"area": t["area"], "rows": rows, "note": note}


def _lines(book, ts, found, names, groups, step_cache):
    """표 묶음 안 수식을 수식 형태로 세어 많이 쓰인 형태부터 줄을 만든다. (줄 목록, 접은 줄 수).

    문장이 같은 수식 형태(RANK 의 순서 인자 0 유무 등)는 줄 하나로 합치고 구역을 모은다.
    풀지 못한 수식의 대체 문장은 방식이 달라도 문장이 같으므로 합치지 않는다."""
    import plain_text
    import tables as tb
    shapes = {}  # 수식 형태 → [(표, 위치, 수식)]
    for t in ts:
        for pos, f in sorted(tb._formula_cells(book, t)):
            shapes.setdefault(plain_text.shape(f, a1(*pos)), []).append((t, pos, f))
    by_text = {}  # 문장 → [줄, 구역 목록, 셀 수]
    for xs in sorted(shapes.values(), key=lambda xs: -len(xs)):
        t, pos, f = xs[0]
        cell = a1(*pos)
        text, example = plain_text.sentence(book, t["sheet"], cell, found, names, groups)
        key = (text, plain_text.shape(f, cell)) if text.endswith(FALLBACK_END) else text
        if key not in by_text:
            by_text[key] = [{"sentence": text, "example": example, "formula": f, "roles": plain_text.line_roles(f),
                             "cell": cell, "step": (t["sheet"], cell)}, [], 0]
        by_text[key][1].extend(_zone(book, x[0]) for x in xs)
        by_text[key][2] += len(xs)
    ranked = sorted(by_text.values(), key=lambda x: -x[2])
    lines = []
    for line, zones, _ in ranked[:LINES_MAX]:
        if line["step"] not in step_cache:
            step_cache[line["step"]] = build(book, *line["step"])
        lines.append({**line, "where": " · ".join(dict.fromkeys(zones))})
    return lines, max(0, len(ranked) - LINES_MAX)


def _card(book, g, ts, found, names, groups, step_cache, row_order=None):
    import plain_text
    first = ts[0]
    kind = "표" if first["header"] else "이름 없는 범위"
    lines, more = _lines(book, ts, found, names, groups, step_cache)
    return {"group": g["id"], "name": names.get(g["id"]) or f"「{first['sheet']}」 {kind}({_area_text(first['area'])})",
            "guessed": g["id"] in names, "text": plain_text.card_text({"sheet": g["sheet"], "tables": ts}),
            "lines": lines, "more": more, "preview": _preview(book, ts), "row_order": row_order}


def _date_note(groups):
    import tables as tb
    if not groups:
        return None
    parts = [("마지막 시트" if g["base"] == max(g["sheets"], key=tb._numbers) else "지정한 시트", g["base"]) for g in groups]
    return " ".join(f"이 화면은 날짜가 붙은 시트 중 {w}({s})를 기준으로 작성했습니다" for w, s in parts)


def build_three_stage(book, analysis, result, names, step_cache):
    """결과표 하나의 페이지 자료. result 는 {"sheet", "excel_table", "row_order"}, names 는 {묶음 번호: 클로드가 유추한 이름},
    step_cache 는 실행 하나에서 공유하는 {(시트, 셀): build 결과} 다."""
    import checks
    import plain_text
    import sources
    import tables as tb
    found, groups, reads = analysis["tables"], analysis["groups"], analysis["reads"]
    result_ids = tb.result_tables(found, groups, reads, result["sheet"], result.get("excel_table"))
    cls = tb.classify(book, found, reads, result_ids)
    rank = {i: n for n, i in enumerate(cls["order"])}
    screens, used = {}, set()
    for key in ("source", "calc", "result"):
        ids = set(cls[key])
        cards = []
        for g in groups:
            ts = [t for t in g["tables"] if t["id"] in ids]
            if ts:
                cards.append((min(rank[t["id"]] for t in ts), g, ts))
                used.update(t["id"] for t in ts)
        screens[key] = [_card(book, g, ts, found, names, groups, step_cache, result["row_order"] if key == "result" else None)
                        for _, g, ts in sorted(cards, key=lambda x: x[0])]
    by_group = {g["id"]: g for g in groups}
    src = [{"group": c["group"], "name": c["name"],
            **sources.source_of(book, [t for t in by_group[c["group"]]["tables"] if t["id"] in cls["source"]])}
           for c in screens["source"]]
    sheets_of = lambda ids: list(dict.fromkeys(t["sheet"] for t in found if t["id"] in ids))  # noqa: E731
    if "same" not in analysis:  # 값이 같은 열은 결과표와 무관하므로 실행 하나에서 1번만 찾는다
        dated = {g["previous"] for g in analysis["date_groups"]}
        by_sheet = {s: [t for t in found if t["sheet"] == s] for s in analysis["sheets"]}
        by_sheet.update({s: tb.find_tables(book, s) for s in dated})
        analysis["same"] = checks.same_value_columns(book, by_sheet, analysis["date_groups"])
    page = [t for t in found if t["id"] in used]
    in_page = lambda o: any(t["sheet"] == o["sheet"] and tb._inside(parse_a1(o["cell"]), t["area"]) for t in page)  # noqa: E731
    return {"title": f"{result.get('excel_table') or result['sheet']} 결과표는 이렇게 만들어집니다",
            "excel": Path(book.path).name, "path": str(book.path), "result": result["sheet"],
            "summary": plain_text.flow_summary(sheets_of(cls["source"]), sheets_of(cls["calc"]), result["sheet"]),
            "date_note": _date_note(analysis["date_groups"]), "screens": screens, "sources": src,
            "source_rules": book.source_rules_state,
            "checks": {"odd": checks.odd_rows(book, page), "same": analysis["same"],
                       "ambiguous": checks.ambiguous(book, [o for o in analysis["odd_refs"] if in_page(o)],
                                                     tb.nameless_read(found, reads, result_ids), analysis["date_groups"])},
            "guide": GUIDE, "steps": list(dict.fromkeys(line["step"] for k in screens for c in screens[k] for line in c["lines"]))}


def render_three_stage(data, step_files):
    """3단계 화면 템플릿에 자료를 넣는다. step_files 는 {(시트, 셀): 단계 화면 파일 이름 또는 None}.
    None 이면 링크 대신 셀 모드 실행 명령을 보인다."""
    links = {}
    for (sheet, cell), name in step_files.items():
        cmd = (f'uv run --system-certs --no-project --with-requirements "{HERE / "requirements.txt"}" python '
               f'"{HERE / "explain_cell.py"}" "{data["path"]}" "{sheet}" {cell} --out explain_{cell}.html')
        links[f"{sheet}!{cell}"] = {"file": name, "command": None if name else cmd}
    page = (HERE.parent / "assets/three_stage_view.html").read_text(encoding="utf-8")
    return page.replace("__DATA__", json.dumps({**data, "links": links}, ensure_ascii=False).replace("<", "\\u003c"))


MAX_OTHER_SHAPES = 4  # 범위 안에서 가장 많은 수식과 모양이 다른 수식을 따라갈 가짓수의 한도
VALUE_RUN_MIN = 6   # 같은 열 · 같은 출처의 값 단계가 이만큼 이어지면 앞 VALUE_RUN_SHOW 개만 보이고 나머지를 접는다
VALUE_RUN_SHOW = 3


def find_repeats(steps, nodes=()):
    """단계 목록에서 접을 구간을 찾는다. 구간마다 `kind` 가 있다.

    `chain`: 같은 열이 위 행 값을 계속 이어 받는 사슬(잔액 · 누적 열). 앞 4단계와 마지막 단계만 남기고 사이를 접는다.
    사슬은 같은 도형의 단계가 바로 앞 같은 도형 단계를 참조하며 이어지는 고리다. 행마다 다른 열 값(입금 등)을
    함께 받으면 단계가 번갈아 나오므로, 접는 범위에는 사이에 끼인 그 입력 단계도 들어간다. 사슬이 6단계 이하면 접지 않는다.
    `values`: 같은 열 · 같은 출처의 값 단계가 6개 이상 이어진 구간. 앞 3개만 남기고 나머지를 접는다.
    """
    out, seen = [], set()
    for i, s in enumerate(steps):
        if i in seen or s["kind"] == "import":
            continue
        chain = [i]
        while True:
            nxt = next((j for j in range(chain[-1] + 1, len(steps))
                        if steps[j]["node"] == s["node"] and chain[-1] in steps[j]["from"]), None)
            if nxt is None:
                break
            chain.append(nxt)
        seen.update(chain)
        if len(chain) > 6:
            out.append({"first": chain[3] + 1, "last": chain[-2], "count": len(chain) - 5,
                        "sheet": steps[chain[4]]["sheet"], "col": steps[chain[4]]["col"],
                        "rows": [steps[chain[4]]["row"], steps[chain[-2]]["row"]], "kind": "chain"})
    hidden = {k for r in out for k in range(r["first"], r["last"] + 1)}
    title = {n["id"]: n["title"] for n in nodes}
    i = 0
    while i < len(steps):
        j = i
        while (steps[i]["kind"] == "value" and i not in hidden and j + 1 < len(steps) and (j + 1) not in hidden
               and steps[j + 1]["kind"] == "value" and steps[j + 1]["node"] == steps[i]["node"]
               and steps[j + 1].get("source") == steps[i].get("source")):
            j += 1
        if j - i + 1 >= VALUE_RUN_MIN:
            fold = range(i + VALUE_RUN_SHOW, j + 1)
            rows = [steps[k]["row"] for k in fold]
            out.append({"first": fold[0], "last": fold[-1], "count": len(fold), "sheet": steps[i]["sheet"],
                        "col": steps[i]["col"], "rows": [min(rows), max(rows)], "name": title.get(steps[i]["node"], ""),
                        "source": steps[i].get("source", ""), "kind": "values"})
        i = j + 1
    return sorted(out, key=lambda r: r["first"])


def name_hints(book, data):
    """도형마다 이름을 정할 재료를 모은다. 스킬을 실행하는 Claude 가 읽고 --names 로 이름을 정한다."""
    first = {}
    for s in data["steps"]:
        if s["kind"] != "import":
            first.setdefault(s["node"], s)
    out = []
    for n in data["nodes"]:
        s = first.get(n["id"])
        if s is None:
            continue
        sh, r, c = s["sheet"], s["row"], s["col"]
        vals = book.sheet_values(sh)
        above = [[a1(rr, c), show(vals[(rr, c)])] for rr in range(r - 1, max(0, r - 31), -1) if (rr, c) in vals][:5]
        left = next(([a1(rr, c - 1), vals[(rr, c - 1)]] for rr in range(r, 0, -1)
                     if isinstance(vals.get((rr, c - 1)), str)), None) if c > 1 else None
        out.append({"id": n["id"], "rule": n["title"], "above": above, "left_label": left,
                    "formula": book.formula(sh, (r, c)), "value": s["value"]})
    return out


SHOWN_ALL, SHOWN_TOP = 20, 10  # 후보 시트가 20개를 넘으면 읽는 표가 많은 10개만 보인다(근거 없는 시작 값)
SIZE_LIMIT = 16 * 1024 * 1024  # 페이지와 단계 화면 크기 합계 상한(근거 없는 시작 값)


def write_candidates(book, analysis, out_dir, base=None):
    """1번째 실행의 후보 파일을 쓰고 경로를 돌려준다. 실행할 때마다 새로 쓴다. base 는 --base 로 받은 {군 이름: 시트 또는 None}."""
    import tables as tb
    by_id = {t["id"]: t for t in analysis["tables"]}
    cands = tb.candidates(analysis["tables"], analysis["reads"])
    top = {c["sheet"] for c in sorted(cands, key=lambda c: -c["reads"])[:SHOWN_TOP]}

    def first_row(t):
        r1, c1, _, c2 = t["area"]
        r = (t["header_row"] or r1 - 1) + 1
        return [show(book.sheet_values(t["sheet"]).get((r, c))) for c in range(c1, c2 + 1)]

    out = {"excel": Path(book.path).name, "base_sheets": base or {},
           "date_groups": [{"pattern": g["pattern"], "count": len(g["sheets"]), "base": g["base"]} for g in analysis["date_groups"]],
           "candidates": [{"sheet": c["sheet"], "reads": c["reads"], "shown": len(cands) <= SHOWN_ALL or c["sheet"] in top,
                           "tables": [{"id": i, "header": list(by_id[i]["header"] or []), "row_order": tb.row_order(book, by_id[i]),
                                       "excel_table": book.tables[by_id[i]["excel_table"]].name if by_id[i]["excel_table"] else None}
                                      for i in c["tables"]]} for c in cands],
           "groups": [{"id": g["id"], "sheet": g["sheet"], "areas": [_area_text(t["area"]) for t in g["tables"]],
                       "header": list(g["tables"][0]["header"] or []), "title": g["tables"][0]["title"],
                       "first_row": first_row(g["tables"][0])} for g in analysis["groups"]]}
    path = Path(out_dir) / "explain_candidates.json"
    path.write_text(json.dumps(out, ensure_ascii=False, indent=1), encoding="utf-8")
    return path


def load_results(path, book, candidates_path):
    """결과 파일을 읽어 검사한다. BOM 을 허용하고 시트 · 기준 시트 이름 앞뒤 공백을 뺀다.
    엑셀 표 이름은 대소문자를 무시하고 찾아 표 사전의 키(대문자)로 바꾼다. 기준 시트는 후보 파일과 같아야 한다."""
    data = json.loads(Path(path).read_text(encoding="utf-8-sig"))
    if data.get("excel") != Path(book.path).name:
        sys.exit(f"결과표 파일의 엑셀 이름 다름: {data.get('excel')}")
    for r in data["results"]:
        r["sheet"] = r["sheet"].strip()
        if r["sheet"] not in book.cells:
            sys.exit(f"결과표 시트 없음: {r['sheet']}")
        table = r.get("excel_table")
        if table:
            key = table.strip().upper()
            if key not in book.tables or book.tables[key].sheet != r["sheet"]:
                sys.exit(f"결과표 엑셀 표 없음: {table}")
            r["excel_table"] = key
    data.setdefault("names", {})
    data["base_sheets"] = {k: v.strip() if isinstance(v, str) else v for k, v in data.get("base_sheets", {}).items()}
    try:
        expected = json.loads(Path(candidates_path).read_text(encoding="utf-8-sig")).get("base_sheets", {})
    except (OSError, ValueError):
        expected = None
    if data["base_sheets"] != expected:
        sys.exit("기준 시트가 후보 파일과 다름(--base 로 1번째 실행부터 다시)")
    return data


def parse_base(values, book):
    """--base 「군 이름=시트」 목록을 {군 이름: 시트 또는 None} 으로 바꾼다. 「=없음」은 군을 푼다."""
    import tables as tb
    groups = {g["pattern"]: g["sheets"] for g in tb.date_groups(book.order)}
    out = {}
    for raw in values or []:
        pattern, _, sheet = raw.partition("=")
        pattern, sheet = pattern.strip(), sheet.strip()
        if pattern not in groups or (sheet != "없음" and sheet not in groups[pattern]):
            sys.exit(f"인자 오류: --base {raw}")
        out[pattern] = None if sheet == "없음" else sheet
    return out


def main_all(a, book):
    """--all 실행. --results 가 없으면 후보 파일만 쓰고, 있으면 결과표마다 페이지와 단계 화면을 쓴다.
    파일은 모두 메모리에서 만든 뒤 출력 경로를 검사하고 쓴다. 중간에 실패하면 아무 파일도 쓰지 않는다."""
    if not any(c.kind == "formula" for cells in book.cells.values() for c in cells.values()):
        sys.exit(f"계산 규칙 없음: {a.excel}")
    out = Path(a.out)
    if not a.results:
        base = parse_base(a.base, book)
        analysis = analyze(book, base)
        path = write_candidates(book, analysis, out.parent, base)
        count = len(json.loads(path.read_text(encoding="utf-8"))["candidates"])
        print(f"결과표 후보 시트 {count}개: {path}")
        return
    results = load_results(a.results, book, out.parent / "explain_candidates.json")
    analysis = analyze(book, results["base_sheets"])
    for r in results["results"]:
        if r["sheet"] in analysis["skipped"]:
            sys.exit(f"결과표 시트가 기준 시트가 아닌 날짜 시트: {r['sheet']}(--base 로 1번째 실행부터 다시)")
    cache, files, total = {}, {}, 0  # files: 경로 → 글자
    for n, result in enumerate(results["results"], 1):
        data = build_three_stage(book, analysis, result, results["names"], cache)
        page = out.parent / f"{out.stem}_r{n}.html"
        steps_dir = out.parent / f"{out.stem}_r{n}_steps"
        step_files = {}
        for i, key in enumerate(data["steps"], 1):
            html = render_page(cache[key])
            size = len(html.encode("utf-8"))
            if total + size > SIZE_LIMIT:
                step_files[key] = None
                continue
            total += size
            files[steps_dir / f"step_{i}.html"] = html
            step_files[key] = f"{steps_dir.name}/step_{i}.html"
        html = render_three_stage(data, step_files)
        total += len(html.encode("utf-8"))
        files[page] = html
    dirs = {p.parent for p in files if p.parent != out.parent}
    pages = [out.parent / f"{out.stem}_r{n}{x}" for n in range(1, len(results["results"]) + 1) for x in (".html", "_steps")]
    if any(p.exists() for p in pages):
        sys.exit("출력 파일 중복(다른 이름 필요)")
    for d in sorted(dirs):
        d.mkdir()
    for p, html in files.items():
        with open(p, "x", encoding="utf-8") as f:
            f.write(html)
    for n in range(1, len(results["results"]) + 1):
        print(out.parent / f"{out.stem}_r{n}.html")


def main(argv=None):
    ap = argparse.ArgumentParser(description="엑셀 셀 하나(또는 --all 로 엑셀 전체)의 계산 과정을 보이는 HTML 화면을 만든다")
    ap.add_argument("excel")
    ap.add_argument("sheet", nargs="?")
    ap.add_argument("cell", nargs="?")
    ap.add_argument("--all", action="store_true", help="시트와 셀 없이 엑셀 전체의 계산 가지를 만든다")
    ap.add_argument("--results", help="--all 일 때 결과표 파일(explain_results.json) 경로. 없으면 후보 파일만 쓴다")
    ap.add_argument("--base", action="append", help="--all 1번째 실행에서 날짜 시트 군의 기준 시트(「군 이름=시트」, 군을 풀면 「군 이름=없음」)")
    ap.add_argument("--out", help="만들 HTML 경로. 이미 있으면 실패한다")
    ap.add_argument("--title", default="", help="결과 셀의 이름(예: 연도별 · 월별 수익률)")
    ap.add_argument("--hints", help="화면 대신 열 이름 재료를 이 JSON 경로에 쓴다")
    ap.add_argument("--names", help="열 이름 JSON {\"시트!열\": \"이름\"} 경로")
    a = ap.parse_args(argv)
    target = a.hints or a.out
    if a.all and a.hints:
        sys.exit("인자 오류: --hints 는 셀 모드 전용")
    if not a.all and not (a.sheet and a.cell):
        sys.exit("인자 누락: 시트와 셀 필요(--all 이 아닐 때)")
    if not target:
        sys.exit("인자 누락: --out 또는 --hints 필요")
    if not a.all and Path(target).exists():  # --all 은 out 이름으로 만든 페이지 경로를 main_all 이 검사한다
        sys.exit(f"출력 파일 중복(다른 이름 필요): {target}")
    names = json.loads(Path(a.names).read_text(encoding="utf-8")) if a.names else None
    if not Path(a.excel).is_file():
        sys.exit(f"엑셀 파일 없음: {a.excel}")
    try:
        book = Book(a.excel)
    except (InvalidFileException, zipfile.BadZipFile):
        sys.exit(f"엑셀 열기 오류(.xlsx · .xlsm 만 지원): {a.excel}")
    if a.all:
        return main_all(a, book)
    try:
        data = build(book, a.sheet, a.cell, a.title, names)
    except ValueError as e:
        sys.exit(str(e))
    if a.hints:
        hints = name_hints(book, data)
        with open(a.hints, "x", encoding="utf-8") as f:  # "x": 그사이 생긴 파일도 덮어쓰지 않는다
            f.write(json.dumps(hints, ensure_ascii=False, indent=1))
        print(a.hints, f"· 열 {len(hints)}개")
        return
    with open(a.out, "x", encoding="utf-8") as f:
        f.write(render_page(data))
    print(a.out, f"· 단계 {len(data['steps'])}개 · 열 {len(data['nodes'])}개")

if __name__ == "__main__":
    main()
